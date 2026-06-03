"""
Bank File Generation Routes for FortexaRH
Generates ACH payment files for Dominican Republic banks (Popular, BHD, Banreservas)
Format verified with real bank templates.
"""
from fastapi import APIRouter, HTTPException, Request, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional
from datetime import datetime, timezone
import io
import logging

router = APIRouter(prefix="/bank-files", tags=["Bank Files"])

from config import db
from utils.auth import get_current_user, get_user_from_request
logger = logging.getLogger(__name__)


class CompanyBankConfig(BaseModel):
    bank_id: str
    account_number: str
    account_type: str = "CC"
    currency: str = "DOP"


# Account type mapping
ACCOUNT_TYPE_MAP = {
    "corriente": "CC",
    "ahorros": "CA",
    "ahorro": "CA",
    "cc": "CC",
    "ca": "CA",
}


def normalize_account_type(account_type: str) -> str:
    """Normalize account type to CC or CA"""
    if not account_type:
        return "CC"
    return ACCOUNT_TYPE_MAP.get(account_type.lower().strip(), "CC")


def format_banreservas_line(
    company_account_type: str, currency: str, company_account: str,
    emp_account_type: str, emp_currency: str, emp_account: str,
    amount: float, concept: str
) -> str:
    """
    Banreservas ACH format (verified from real bank template):
    CC,DOP,0130850482,CC,DOP,9608649339,9409.00,AUXILIAR DE CONTABILIDAD
    Fields: TipoCuentaEmpresa,Moneda,CuentaEmpresa,TipoCuentaEmpleado,Moneda,CuentaEmpleado,Monto,Concepto
    """
    concept_clean = concept.upper().replace(",", " ").strip()[:50]
    return f"{company_account_type},{currency},{company_account},{emp_account_type},{emp_currency},{emp_account},{amount:.2f},{concept_clean}\n"


def format_popular_line(seq: int, account: str, amount: float, name: str, doc_type: str, doc_number: str) -> str:
    """Format a line for Banco Popular file"""
    return f"{seq:06d}|22|{account}|{amount:.2f}|{name[:40]}|{doc_type}|{doc_number}|\n"


def format_bhd_line(account: str, amount: float, name: str, doc_number: str) -> str:
    """Format a line for BHD León file"""
    return f"{account:<20}{amount:>15.2f}{name:<40}{doc_number:<15}\n"


@router.get("/banks")
async def get_available_banks(request: Request):
    """Get list of available banks for file generation"""
    await get_user_from_request(request)
    return [
        {"id": "banreservas", "name": "Banreservas", "format": "TXT", "description": "Formato ACH Banreservas (TXT delimitado o Excel oficial)", "formats": ["txt", "xlsx"]},
        {"id": "popular", "name": "Banco Popular Dominicano", "format": "TXT", "description": "Formato Nómina Popular", "formats": ["txt"]},
        {"id": "bhd", "name": "BHD León", "format": "TXT", "description": "Formato ACH BHD", "formats": ["txt"]},
    ]


@router.get("/company-bank-config")
async def get_company_bank_config(request: Request, bank_id: Optional[str] = None):
    """Get company bank configuration for ACH generation"""
    current_user = await get_user_from_request(request)
    company_id = current_user.get("company_id")
    
    if bank_id:
        config = await db.company_bank_config.find_one(
            {"company_id": company_id, "bank_id": bank_id},
            {"_id": 0}
        )
        return config or {"company_id": company_id, "bank_id": bank_id}
    
    configs = await db.company_bank_config.find(
        {"company_id": company_id},
        {"_id": 0}
    ).to_list(10)
    
    return {"company_id": company_id, "accounts": configs}


@router.put("/company-bank-config")
async def save_company_bank_config(data: CompanyBankConfig, request: Request):
    """Save company bank account for ACH file generation"""
    current_user = await get_user_from_request(request)
    company_id = current_user.get("company_id")
    
    await db.company_bank_config.update_one(
        {"company_id": company_id, "bank_id": data.bank_id},
        {"$set": {
            "company_id": company_id,
            "bank_id": data.bank_id,
            "account_number": data.account_number,
            "account_type": data.account_type,
            "currency": data.currency,
            "updated_at": datetime.now(timezone.utc).isoformat()
        }},
        upsert=True
    )
    
    return {"success": True, "message": "Configuración bancaria guardada"}


@router.get("/generate/{period_id}/{bank_id}")
async def generate_bank_file(period_id: str, bank_id: str, request: Request, format: str = "txt"):
    """Generate ACH bank payment file for a payroll period.

    For Banreservas the caller may choose `format=txt` (default, official ACH
    delimited file) or `format=xlsx` (official Banreservas Nómina Electrónica
    Excel template, ready to upload via the bank portal)."""
    current_user = await get_user_from_request(request)
    company_id = current_user.get("company_id")
    
    # Get payroll period
    period = await db.payroll_periods.find_one(
        {"period_id": period_id, "company_id": company_id},
        {"_id": 0}
    )
    
    if not period:
        raise HTTPException(status_code=404, detail="Período no encontrado")
    
    if period.get("status") not in ["approved", "paid", "processed"]:
        raise HTTPException(status_code=400, detail="El período debe estar aprobado o pagado para generar archivo ACH")
    
    # Get company bank config for this bank
    company_bank = await db.company_bank_config.find_one(
        {"company_id": company_id, "bank_id": bank_id},
        {"_id": 0}
    )
    
    company_account = company_bank.get("account_number", "") if company_bank else ""
    company_account_type = company_bank.get("account_type", "CC") if company_bank else "CC"
    company_currency = company_bank.get("currency", "DOP") if company_bank else "DOP"
    
    # Get payroll entries
    payrolls = await db.payroll_entries.find(
        {"period_id": period_id, "company_id": company_id},
        {"_id": 0}
    ).to_list(1000)
    
    if not payrolls:
        raise HTTPException(status_code=400, detail="No hay nóminas en este período")
    
    # Get employee bank info
    employee_ids = [p.get("employee_id") for p in payrolls]
    employees = await db.employees.find(
        {"employee_id": {"$in": employee_ids}, "company_id": company_id},
        {"_id": 0, "employee_id": 1, "first_name": 1, "last_name": 1,
         "document_number": 1, "account_number": 1, "account_type": 1, "bank_name": 1, "position": 1}
    ).to_list(1000)
    
    emp_map = {e["employee_id"]: e for e in employees}
    
    # Track employees without bank info
    missing_bank = []
    
    output = io.StringIO()
    record_count = 0
    total_amount = 0.0
    
    if bank_id == "banreservas":
        # Banreservas: TXT (CSV-delimited ACH) or official XLSX template
        bank_format = (format or "txt").lower()
        if bank_format not in ("txt", "xlsx"):
            bank_format = "txt"

        # Map TXT codes (CC/CA) to human labels expected by the Excel template
        type_label = {"CC": "Corriente", "CA": "Ahorro"}
        currency_label = {"DOP": "Pesos", "USD": "Dólares"}
        company_type_label = type_label.get(company_account_type, "Corriente")
        currency_label_text = currency_label.get(company_currency, "Pesos")

        rows: list[dict] = []
        for payroll in payrolls:
            emp = emp_map.get(payroll.get("employee_id"), {})
            emp_account = emp.get("account_number", "")
            emp_account_type = normalize_account_type(emp.get("account_type", "CC"))
            amount = payroll.get("net_salary", 0)
            concept = emp.get("position", payroll.get("position", "PAGO NOMINA"))

            if not emp_account:
                missing_bank.append(f"{emp.get('first_name', '')} {emp.get('last_name', '')}")
                continue
            if amount <= 0:
                continue

            full_name = f"{emp.get('last_name', '').upper()}, {emp.get('first_name', '').upper()}".strip(", ")
            concept_clean = (concept or "PAGO NOMINA").upper().replace(",", " ").strip()[:50]

            rows.append({
                "name": full_name,
                "company_account_type": company_account_type,
                "company_type_label": company_type_label,
                "currency": company_currency,
                "currency_label": currency_label_text,
                "company_account": company_account,
                "emp_account_type": emp_account_type,
                "emp_type_label": type_label.get(emp_account_type, "Corriente"),
                "emp_account": emp_account,
                "amount": amount,
                "concept": concept_clean,
            })
            record_count += 1
            total_amount += amount

        if bank_format == "xlsx":
            # Build the official Banreservas Excel template
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = Workbook()
            ws = wb.active
            ws.title = "CONVERTIR NOMINA"

            label_font = Font(bold=True, size=11)
            value_font = Font(size=11)
            header_fill = PatternFill("solid", fgColor="1F4E78")
            header_font = Font(bold=True, color="FFFFFF", size=11)
            thin = Side(border_style="thin", color="B7B7B7")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            meta = [
                ("TIPO CUENTA EMPRESA", company_type_label),
                ("NUMERO DE CUENTA EMPRESA", company_account),
                ("MONEDA A PAGAR", currency_label_text),
                ("TOTAL DE NÓMINA", round(total_amount, 2)),
                ("CANTIDAD DE EMPLEADOS", record_count),
            ]
            for i, (label, value) in enumerate(meta, 1):
                ws.cell(row=i, column=1, value=label).font = label_font
                ws.cell(row=i, column=5, value=value).font = value_font

            headers = [
                "Nombre de empleado", "TIPO CUENTA EMPRESA", "TIPO DE LA MONEDA",
                "NUMERO DE CUENTA DE EMPRESA", "Tipo cuenta empleado", "MONEDA A PAGAR",
                "Numero cuenta empleado", "Monto a pagar", "Concepto",
            ]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=8, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border

            for idx, r in enumerate(rows, start=9):
                values = [
                    r["name"], r["company_type_label"], r["currency_label"],
                    r["company_account"], r["emp_type_label"], r["currency_label"],
                    r["emp_account"], round(float(r["amount"]), 2), r["concept"],
                ]
                for col, v in enumerate(values, 1):
                    cell = ws.cell(row=idx, column=col, value=v)
                    cell.border = border
                    cell.font = value_font
                    if col == 8:
                        cell.number_format = '#,##0.00'

            for col, width in enumerate([34, 20, 18, 24, 20, 16, 22, 16, 32], 1):
                ws.column_dimensions[chr(64 + col)].width = width
            ws.freeze_panes = "A9"

            xlsx_buf = io.BytesIO()
            wb.save(xlsx_buf)
            xlsx_bytes = xlsx_buf.getvalue()
            xlsx_buf.close()

            filename = f"Nomina_Banreservas_{period.get('description', period_id).replace(' ', '_')}.xlsx"
            await db.bank_file_logs.insert_one({
                "company_id": company_id, "period_id": period_id, "bank_id": bank_id,
                "format": "xlsx", "record_count": record_count,
                "missing_bank_info": missing_bank, "total_amount": round(total_amount, 2),
                "generated_by": current_user.get("user_id"),
                "generated_at": datetime.now(timezone.utc).isoformat(), "filename": filename,
            })
            if record_count == 0:
                raise HTTPException(
                    status_code=400,
                    detail=f"No se generaron registros. {len(missing_bank)} empleados sin datos bancarios: {', '.join(missing_bank[:5])}",
                )
            return StreamingResponse(
                io.BytesIO(xlsx_bytes),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={
                    "Content-Disposition": f"attachment; filename={filename}",
                    "X-Record-Count": str(record_count),
                    "X-Total-Amount": f"{total_amount:.2f}",
                    "X-Missing-Bank": str(len(missing_bank)),
                },
            )

        # Default: TXT format
        for r in rows:
            output.write(format_banreservas_line(
                r["company_account_type"], r["currency"], r["company_account"],
                r["emp_account_type"], r["currency"], r["emp_account"],
                r["amount"], r["concept"],
            ))

        filename = f"ACH_Banreservas_{period.get('description', period_id).replace(' ', '_')}.txt"
        
    elif bank_id == "popular":
        output.write(f"H|{company_account}|{datetime.now().strftime('%Y%m%d')}|NOMINA\n")
        for i, payroll in enumerate(payrolls, 1):
            emp = emp_map.get(payroll.get("employee_id"), {})
            account = emp.get("account_number", "")
            name = f"{emp.get('last_name', '')},{emp.get('first_name', '')}".strip()
            doc = emp.get("document_number", "")
            amount = payroll.get("net_salary", 0)
            
            if not account:
                missing_bank.append(f"{emp.get('first_name', '')} {emp.get('last_name', '')}")
                continue
            
            if amount > 0:
                output.write(format_popular_line(i, account, amount, name, "C", doc))
                record_count += 1
                total_amount += amount
        
        output.write(f"T|{record_count}|{total_amount:.2f}\n")
        filename = f"ACH_Popular_{period.get('description', period_id).replace(' ', '_')}.txt"
        
    elif bank_id == "bhd":
        for payroll in payrolls:
            emp = emp_map.get(payroll.get("employee_id"), {})
            account = emp.get("account_number", "")
            name = f"{emp.get('last_name', '')},{emp.get('first_name', '')}".strip()
            doc = emp.get("document_number", "")
            amount = payroll.get("net_salary", 0)
            
            if not account:
                missing_bank.append(f"{emp.get('first_name', '')} {emp.get('last_name', '')}")
                continue
            
            if amount > 0:
                output.write(format_bhd_line(account, amount, name, doc))
                record_count += 1
                total_amount += amount
        
        filename = f"ACH_BHD_{period.get('description', period_id).replace(' ', '_')}.txt"
    else:
        raise HTTPException(status_code=400, detail="Banco no soportado")
    
    # Log the generation
    await db.bank_file_logs.insert_one({
        "company_id": company_id,
        "period_id": period_id,
        "bank_id": bank_id,
        "record_count": record_count,
        "missing_bank_info": missing_bank,
        "total_amount": round(total_amount, 2),
        "generated_by": current_user.get("user_id"),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "filename": filename
    })
    
    content = output.getvalue()
    output.close()
    
    if not content.strip():
        raise HTTPException(
            status_code=400,
            detail=f"No se generaron registros. {len(missing_bank)} empleados sin datos bancarios: {', '.join(missing_bank[:5])}"
        )
    
    return StreamingResponse(
        io.BytesIO(content.encode('utf-8')),
        media_type="text/plain; charset=utf-8",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "X-Record-Count": str(record_count),
            "X-Total-Amount": f"{total_amount:.2f}",
            "X-Missing-Bank": str(len(missing_bank))
        }
    )


@router.get("/preview/{period_id}/{bank_id}")
async def preview_bank_file(period_id: str, bank_id: str, request: Request):
    """Preview bank file data before generating (shows which employees are ready and which are missing bank info)"""
    current_user = await get_user_from_request(request)
    company_id = current_user.get("company_id")
    
    period = await db.payroll_periods.find_one(
        {"period_id": period_id, "company_id": company_id},
        {"_id": 0}
    )
    if not period:
        raise HTTPException(status_code=404, detail="Período no encontrado")
    
    payrolls = await db.payroll_entries.find(
        {"period_id": period_id, "company_id": company_id},
        {"_id": 0}
    ).to_list(1000)
    
    employee_ids = [p.get("employee_id") for p in payrolls]
    employees = await db.employees.find(
        {"employee_id": {"$in": employee_ids}, "company_id": company_id},
        {"_id": 0, "employee_id": 1, "first_name": 1, "last_name": 1,
         "account_number": 1, "account_type": 1, "bank_name": 1, "position": 1}
    ).to_list(1000)
    
    emp_map = {e["employee_id"]: e for e in employees}
    
    ready = []
    missing = []
    total_amount = 0.0
    
    for payroll in payrolls:
        emp = emp_map.get(payroll.get("employee_id"), {})
        name = f"{emp.get('first_name', '')} {emp.get('last_name', '')}".strip()
        amount = payroll.get("net_salary", 0)
        
        if emp.get("account_number"):
            ready.append({
                "employee_name": name,
                "account": emp.get("account_number"),
                "account_type": emp.get("account_type", "Corriente"),
                "amount": round(amount, 2),
                "position": emp.get("position", "")
            })
            total_amount += amount
        else:
            missing.append({
                "employee_name": name,
                "employee_id": payroll.get("employee_id"),
                "amount": round(amount, 2)
            })
    
    company_bank = await db.company_bank_config.find_one(
        {"company_id": company_id, "bank_id": bank_id},
        {"_id": 0}
    )
    
    return {
        "period": period.get("description", period_id),
        "bank_id": bank_id,
        "company_account": company_bank.get("account_number", "") if company_bank else "",
        "ready_count": len(ready),
        "missing_count": len(missing),
        "total_amount": round(total_amount, 2),
        "ready": ready,
        "missing": missing
    }


@router.get("/history")
async def get_bank_file_history(request: Request):
    """Get history of generated bank files"""
    current_user = await get_user_from_request(request)
    company_id = current_user.get("company_id")
    
    history = await db.bank_file_logs.find(
        {"company_id": company_id},
        {"_id": 0}
    ).sort("generated_at", -1).to_list(50)
    
    return history
