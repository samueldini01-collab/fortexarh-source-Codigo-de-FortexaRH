"""
Employee Loans Routes for FortexaRH
Module for managing employee loans, payments, and deductions
"""
from fastapi import APIRouter, HTTPException, Depends, Request
from pydantic import BaseModel
from typing import Optional, List
from datetime import datetime, timezone, timedelta
import uuid
import logging

router = APIRouter(prefix="/loans", tags=["Employee Loans"])

from config import db
from utils.auth import get_current_user, get_user_from_request
logger = logging.getLogger(__name__)


# ===================== PYDANTIC MODELS =====================
from models.finance import LoanCreate, LoanPaymentCreate


# ===================== ROUTER INITIALIZATION =====================


# ===================== LOAN ENDPOINTS =====================

@router.get("")
async def get_loans(
    request: Request,
    status: Optional[str] = None,
    employee_id: Optional[str] = None
):
    """Get all loans for the company"""
    current_user = await get_user_from_request(request)
    company_id = current_user.get("company_id")
    
    query = {"company_id": company_id}
    if status:
        query["status"] = status
    if employee_id:
        query["employee_id"] = employee_id
    
    loans = await db.loans.find(query, {"_id": 0}).sort("created_at", -1).to_list(500)
    
    # Enrich with employee info
    for loan in loans:
        employee = await db.employees.find_one(
            {"employee_id": loan.get("employee_id"), "company_id": company_id},
            {"_id": 0, "first_name": 1, "last_name": 1, "document_number": 1, "position": 1}
        )
        if employee:
            loan["employee_name"] = f"{employee.get('first_name', '')} {employee.get('last_name', '')}"
            loan["employee_document"] = employee.get("document_number", "")
            loan["employee_position"] = employee.get("position", "")
    
    return loans


@router.get("/summary")
async def get_loans_summary(request: Request):
    """Get loans summary for dashboard.

    Returns full KPIs across all statuses (active/paid/paused/cancelled)
    plus the active aggregate amounts used by the ``/loans`` dashboard.
    """
    current_user = await get_user_from_request(request)
    company_id = current_user.get("company_id")

    all_loans = await db.loans.find({"company_id": company_id}, {"_id": 0}).to_list(2000)
    active_loans = [l for l in all_loans if l.get("status") == "active"]
    paid_loans = [l for l in all_loans if l.get("status") in ("paid", "paid_off")]
    paused_loans = [l for l in all_loans if l.get("status") == "paused"]
    cancelled_loans = [l for l in all_loans if l.get("status") == "cancelled"]

    # Active = what's still owed by employees
    total_pending = sum(loan.get("remaining_balance", 0) for loan in active_loans)
    # Total paid across entire history (any status)
    total_paid = sum(loan.get("total_paid", 0) for loan in all_loans)
    # Total ever loaned (capital, any status)
    total_loaned = sum(loan.get("amount", 0) for loan in all_loans)
    # What gets withheld this month from payrolls (active only)
    monthly_total = sum(loan.get("monthly_payment", 0) for loan in active_loans)

    status_counts = {}
    for loan in all_loans:
        status = loan.get("status", "unknown")
        status_counts[status] = status_counts.get(status, 0) + 1

    return {
        "total_active_loans": len(active_loans),
        "total_paid_loans": len(paid_loans),
        "total_paused_loans": len(paused_loans),
        "total_cancelled_loans": len(cancelled_loans),
        "total_loans": len(all_loans),
        "total_loaned": round(total_loaned, 2),
        "total_paid": round(total_paid, 2),
        "total_pending": round(total_pending, 2),
        "monthly_deduction_total": round(monthly_total, 2),
        "employees_with_loans": len({loan.get("employee_id") for loan in active_loans}),
        "status_counts": status_counts,
    }


@router.get("/export/xlsx")
async def export_loans_xlsx(
    request: Request,
    status: Optional[str] = None,
    employee_id: Optional[str] = None,
):
    """Export loans to a styled XLSX file with KPIs and a totals row."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse

    current_user = await get_user_from_request(request)
    company_id = current_user.get("company_id")

    query = {"company_id": company_id}
    if status:
        query["status"] = status
    if employee_id:
        query["employee_id"] = employee_id

    loans = await db.loans.find(query, {"_id": 0}).sort("created_at", -1).to_list(2000)

    emp_map = {}
    if loans:
        emp_ids = list({l.get("employee_id") for l in loans})
        emps = await db.employees.find(
            {"employee_id": {"$in": emp_ids}, "company_id": company_id},
            {"_id": 0, "employee_id": 1, "first_name": 1, "last_name": 1, "document_number": 1},
        ).to_list(2000)
        for e in emps:
            emp_map[e["employee_id"]] = (
                f"{e.get('first_name', '')} {e.get('last_name', '')}".strip(),
                e.get("document_number", ""),
            )

    company = await db.companies.find_one({"company_id": company_id}, {"_id": 0, "name": 1})
    company_name = (company or {}).get("name", "FortexaRH")

    STATUS_LABEL = {"active": "Activo", "paused": "Pausado", "paid": "Pagado",
                    "paid_off": "Pagado", "cancelled": "Cancelado"}
    METHOD_LABEL = {"linear": "Lineal", "french": "Francesa", None: "Sin interés", "": "Sin interés"}
    SCHEDULE_LABEL = {"all_periods": "Todos los períodos",
                      "monthly_only": "Solo mensual",
                      "biweekly_second_only": "Solo 2da quincena"}

    wb = Workbook()
    ws = wb.active
    ws.title = "Préstamos"

    title_font = Font(bold=True, size=14, color="FFFFFF")
    title_fill = PatternFill("solid", fgColor="1E293B")
    header_font = Font(bold=True, size=10, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="334155")
    totals_font = Font(bold=True, size=10)
    totals_fill = PatternFill("solid", fgColor="E2E8F0")
    kpi_fill = PatternFill("solid", fgColor="F1F5F9")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    thin = Side(border_style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    columns = [
        ("id", "ID Préstamo", 16, "left"),
        ("employee_name", "Empleado", 28, "left"),
        ("document", "Cédula", 14, "left"),
        ("description", "Descripción", 24, "left"),
        ("amount", "Capital", 14, "right"),
        ("interest_rate", "Tasa %", 8, "right"),
        ("interest_method", "Método", 12, "left"),
        ("schedule_type", "Calendario", 18, "left"),
        ("term_months", "Plazo (m)", 9, "right"),
        ("monthly_payment", "Cuota Mensual", 14, "right"),
        ("total_to_pay", "Total a Pagar", 14, "right"),
        ("total_paid", "Total Pagado", 14, "right"),
        ("remaining_balance", "Saldo Pendiente", 16, "right"),
        ("status", "Estado", 11, "center"),
        ("start_date", "Inicio", 12, "center"),
    ]
    n_cols = len(columns)
    currency_fmt = '_-#,##0.00_-;[Red]-#,##0.00_-'

    ws.cell(row=1, column=1, value=f"{company_name} — Préstamos a Empleados").font = title_font
    ws.cell(row=1, column=1).fill = title_fill
    ws.cell(row=1, column=1).alignment = center
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    ws.row_dimensions[1].height = 24

    filter_parts = []
    if status:
        filter_parts.append(f"Estado: {STATUS_LABEL.get(status, status)}")
    if employee_id:
        nm = emp_map.get(employee_id, ("", ""))[0]
        filter_parts.append(f"Empleado: {nm or employee_id}")
    ws.cell(row=2, column=1, value=" · ".join(filter_parts) or "Todos los préstamos").font = Font(italic=True, size=10, color="475569")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)

    total_loaned = sum(l.get("amount", 0) for l in loans)
    total_paid = sum(l.get("total_paid", 0) for l in loans)
    total_remaining = sum(l.get("remaining_balance", 0) for l in loans if l.get("status") == "active")
    n_active = sum(1 for l in loans if l.get("status") == "active")
    n_emps = len({l.get("employee_id") for l in loans if l.get("status") == "active"})
    kpi_text = (
        f"  KPIs:  Préstamos: {len(loans)}  ·  Activos: {n_active}  ·  "
        f"Empleados con préstamos activos: {n_emps}  ·  "
        f"Total prestado: RD${total_loaned:,.2f}  ·  "
        f"Total cobrado: RD${total_paid:,.2f}  ·  "
        f"Pendiente: RD${total_remaining:,.2f}"
    )
    ws.cell(row=3, column=1, value=kpi_text).font = Font(bold=True, size=10, color="1E293B")
    ws.cell(row=3, column=1).fill = kpi_fill
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=n_cols)
    ws.row_dimensions[3].height = 22

    HEADER_ROW = 5
    for c_idx, (_, label, _w, _a) in enumerate(columns, start=1):
        cell = ws.cell(row=HEADER_ROW, column=c_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[HEADER_ROW].height = 30
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=3)

    for r_idx, loan in enumerate(loans, start=HEADER_ROW + 1):
        emp_name, emp_doc = emp_map.get(loan.get("employee_id"), ("", ""))
        row_values = {
            "id": loan.get("loan_id", ""),
            "employee_name": emp_name or "",
            "document": emp_doc or "",
            "description": loan.get("description", "") or "",
            "amount": float(loan.get("amount", 0) or 0),
            "interest_rate": float(loan.get("interest_rate", 0) or 0),
            "interest_method": METHOD_LABEL.get(loan.get("interest_method"), loan.get("interest_method") or "—"),
            "schedule_type": SCHEDULE_LABEL.get(loan.get("schedule_type"), loan.get("schedule_type") or "—"),
            "term_months": int(loan.get("term_months", 0) or 0),
            "monthly_payment": float(loan.get("monthly_payment", 0) or 0),
            "total_to_pay": float(loan.get("total_to_pay", 0) or 0),
            "total_paid": float(loan.get("total_paid", 0) or 0),
            "remaining_balance": float(loan.get("remaining_balance", 0) or 0),
            "status": STATUS_LABEL.get(loan.get("status"), loan.get("status") or ""),
            "start_date": loan.get("start_date", "") or "",
        }
        for c_idx, (key, _label, _w, align) in enumerate(columns, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=row_values[key])
            cell.border = border
            cell.alignment = right if align == "right" else center if align == "center" else left
            if key in {"amount", "monthly_payment", "total_to_pay", "total_paid", "remaining_balance"}:
                cell.number_format = currency_fmt

    totals_row = HEADER_ROW + 1 + len(loans)
    totals_keys = {
        "amount": total_loaned,
        "total_paid": total_paid,
        "remaining_balance": total_remaining,
        "monthly_payment": sum(l.get("monthly_payment", 0) for l in loans if l.get("status") == "active"),
        "total_to_pay": sum(l.get("total_to_pay", 0) for l in loans),
    }
    for c_idx, (key, _label, _w, _align) in enumerate(columns, start=1):
        cell = ws.cell(row=totals_row, column=c_idx)
        cell.font = totals_font
        cell.fill = totals_fill
        cell.border = border
        if c_idx == 1:
            cell.value = "TOTALES"
            cell.alignment = center
        elif key in totals_keys:
            cell.value = totals_keys[key]
            cell.alignment = right
            cell.number_format = currency_fmt

    for c_idx, (_k, _l, w, _a) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"prestamos_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-cache"},
    )


@router.post("")
async def create_loan(data: LoanCreate, request: Request):
    """Create a new loan for an employee"""
    current_user = await get_user_from_request(request)
    company_id = current_user.get("company_id")
    
    # Verify employee exists
    employee = await db.employees.find_one(
        {"employee_id": data.employee_id, "company_id": company_id},
        {"_id": 0}
    )
    if not employee:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")

    # Calculate monthly payment based on interest method
    if data.interest_rate > 0 and data.interest_method == "french":
        # French amortization (interest on outstanding balance, cuota fija)
        monthly_rate = data.interest_rate / 100 / 12
        monthly_payment = data.amount * (monthly_rate * (1 + monthly_rate) ** data.term_months) / ((1 + monthly_rate) ** data.term_months - 1)
        total_to_pay = monthly_payment * data.term_months
    elif data.interest_rate > 0:
        # Linear simple interest: total = capital + (capital × rate × months/12)
        # Cuota = total / num_cuotas (flat, no recalculation)
        total_interest = data.amount * (data.interest_rate / 100) * (data.term_months / 12)
        total_to_pay = data.amount + total_interest
        monthly_payment = total_to_pay / data.term_months
    else:
        # No interest (simple division)
        monthly_payment = data.amount / data.term_months
        total_to_pay = data.amount

    loan_id = f"loan_{uuid.uuid4().hex[:12]}"

    # Generate payment schedule
    schedule = []
    start = datetime.strptime(data.start_date, "%Y-%m-%d")
    remaining = data.amount

    for i in range(data.term_months):
        payment_date = start + timedelta(days=30 * (i + 1))

        if data.interest_rate > 0 and data.interest_method == "french":
            interest_payment = remaining * (data.interest_rate / 100 / 12)
            principal_payment = monthly_payment - interest_payment
        elif data.interest_rate > 0:
            # Linear: each installment carries equal principal + equal interest
            interest_payment = (data.amount * (data.interest_rate / 100) * (data.term_months / 12)) / data.term_months
            principal_payment = data.amount / data.term_months
        else:
            interest_payment = 0
            principal_payment = monthly_payment

        remaining = max(0, remaining - principal_payment)

        schedule.append({
            "installment_number": i + 1,
            "due_date": payment_date.strftime("%Y-%m-%d"),
            "amount": round(monthly_payment, 2),
            "principal": round(principal_payment, 2),
            "interest": round(interest_payment, 2),
            "remaining_balance": round(remaining, 2),
            "status": "pending",
            "paid_date": None,
            "paid_amount": 0
        })

    loan = {
        "loan_id": loan_id,
        "company_id": company_id,
        "employee_id": data.employee_id,
        "amount": data.amount,
        "currency": data.currency,
        "interest_rate": data.interest_rate,
        "interest_method": data.interest_method,
        "schedule_type": data.schedule_type,
        "term_months": data.term_months,
        "monthly_payment": round(monthly_payment, 2),
        "total_to_pay": round(total_to_pay, 2),
        "total_paid": 0,
        "remaining_balance": round(total_to_pay, 2),
        "start_date": data.start_date,
        "description": data.description,
        "deduct_from_payroll": data.deduct_from_payroll,
        "status": "active",
        "payment_schedule": schedule,
        "payments": [],
        "created_by": current_user.get("user_id"),
        "created_at": datetime.now(timezone.utc).isoformat(),
        "updated_at": datetime.now(timezone.utc).isoformat()
    }

    await db.loans.insert_one(loan)

    return {
        "loan_id": loan_id,
        "monthly_payment": round(monthly_payment, 2),
        "total_to_pay": round(total_to_pay, 2),
        "message": "Préstamo creado exitosamente"
    }


# ===================== STATUS TRANSITIONS =====================


async def _set_loan_status(loan_id: str, company_id: str, new_status: str, expected_from: list[str]):
    loan = await db.loans.find_one(
        {"loan_id": loan_id, "company_id": company_id},
        {"_id": 0}
    )
    if not loan:
        raise HTTPException(status_code=404, detail="Préstamo no encontrado")
    if loan.get("status") not in expected_from:
        raise HTTPException(
            status_code=400,
            detail=f"Transición inválida: estado actual '{loan.get('status')}', esperado {expected_from}"
        )
    await db.loans.update_one(
        {"loan_id": loan_id, "company_id": company_id},
        {"$set": {"status": new_status, "updated_at": datetime.now(timezone.utc).isoformat()}}
    )
    return {"loan_id": loan_id, "status": new_status, "message": f"Préstamo {new_status}"}


@router.post("/{loan_id}/pause")
async def pause_loan(loan_id: str, request: Request):
    """Pause an active loan — no automatic deductions until resumed."""
    current_user = await get_user_from_request(request)
    return await _set_loan_status(loan_id, current_user.get("company_id"), "paused", ["active"])


@router.post("/{loan_id}/resume")
async def resume_loan(loan_id: str, request: Request):
    """Resume a paused loan."""
    current_user = await get_user_from_request(request)
    return await _set_loan_status(loan_id, current_user.get("company_id"), "active", ["paused"])


@router.post("/{loan_id}/cancel")
async def cancel_loan(loan_id: str, request: Request):
    """Cancel a loan (forfeit remaining balance)."""
    current_user = await get_user_from_request(request)
    return await _set_loan_status(loan_id, current_user.get("company_id"), "cancelled", ["active", "paused"])


@router.get("/{loan_id}")
async def get_loan(loan_id: str, request: Request):
    """Get loan details"""
    current_user = await get_user_from_request(request)
    company_id = current_user.get("company_id")
    
    loan = await db.loans.find_one(
        {"loan_id": loan_id, "company_id": company_id},
        {"_id": 0}
    )
    
    if not loan:
        raise HTTPException(status_code=404, detail="Préstamo no encontrado")
    
    # Get employee info
    employee = await db.employees.find_one(
        {"employee_id": loan.get("employee_id"), "company_id": company_id},
        {"_id": 0, "first_name": 1, "last_name": 1, "document_number": 1, "position": 1, "department": 1}
    )
    
    if employee:
        loan["employee"] = {
            "name": f"{employee.get('first_name', '')} {employee.get('last_name', '')}",
            "document": employee.get("document_number", ""),
            "position": employee.get("position", ""),
            "department": employee.get("department", "")
        }
    
    return loan


@router.post("/{loan_id}/payment")
async def register_loan_payment(loan_id: str, data: LoanPaymentCreate, request: Request):
    """Register a payment for a loan"""
    current_user = await get_user_from_request(request)
    company_id = current_user.get("company_id")
    
    loan = await db.loans.find_one(
        {"loan_id": loan_id, "company_id": company_id},
        {"_id": 0}
    )
    
    if not loan:
        raise HTTPException(status_code=404, detail="Préstamo no encontrado")
    
    if loan.get("status") != "active":
        raise HTTPException(status_code=400, detail="El préstamo no está activo")
    
    payment_id = f"pay_{uuid.uuid4().hex[:8]}"
    new_total_paid = loan.get("total_paid", 0) + data.amount
    new_remaining = loan.get("amount", 0) - new_total_paid
    
    payment = {
        "payment_id": payment_id,
        "amount": data.amount,
        "payment_date": data.payment_date,
        "payment_type": data.payment_type,
        "notes": data.notes,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    
    # Update next pending installment
    schedule = loan.get("payment_schedule", [])
    for installment in schedule:
        if installment.get("status") == "pending":
            installment["status"] = "paid"
            installment["paid_date"] = data.payment_date
            installment["paid_amount"] = data.amount
            break
    
    # Check if loan is fully paid
    new_status = "paid" if new_remaining <= 0 else "active"
    
    await db.loans.update_one(
        {"loan_id": loan_id, "company_id": company_id},
        {
            "$set": {
                "total_paid": round(new_total_paid, 2),
                "remaining_balance": round(max(0, new_remaining), 2),
                "status": new_status,
                "payment_schedule": schedule,
                "updated_at": datetime.now(timezone.utc).isoformat()
            },
            "$push": {"payments": payment}
        }
    )
    
    return {
        "payment_id": payment_id,
        "total_paid": round(new_total_paid, 2),
        "remaining_balance": round(max(0, new_remaining), 2),
        "status": new_status,
        "message": "Pago registrado exitosamente"
    }


@router.delete("/{loan_id}")
async def delete_loan(loan_id: str, request: Request):
    """Delete a loan (only if no payments made)"""
    current_user = await get_user_from_request(request)
    company_id = current_user.get("company_id")
    
    loan = await db.loans.find_one(
        {"loan_id": loan_id, "company_id": company_id},
        {"_id": 0}
    )
    
    if not loan:
        raise HTTPException(status_code=404, detail="Préstamo no encontrado")
    
    if loan.get("total_paid", 0) > 0:
        raise HTTPException(status_code=400, detail="No se puede eliminar un préstamo con pagos registrados")
    
    await db.loans.delete_one({"loan_id": loan_id, "company_id": company_id})
    
    return {"message": "Préstamo eliminado"}


# ===================== PAYROLL AUTO-DEDUCTION HELPERS =====================


def loan_deduction_for_period(loan: dict, period_type: str | None) -> float:
    """Return the deduction amount this single payroll period should withhold
    from the given active loan, taking into account:
      * ``schedule_type`` configured on the loan
      * ``period_type`` (mensual / quincenal_1 / quincenal_2)
      * Cap at ``remaining_balance`` so we never over-collect
    """
    if loan.get("status") != "active" or not loan.get("deduct_from_payroll", True):
        return 0.0
    monthly = float(loan.get("monthly_payment", 0) or 0)
    remaining = float(loan.get("remaining_balance", 0) or 0)
    if monthly <= 0 or remaining <= 0:
        return 0.0

    pt = (period_type or "mensual")
    sched = loan.get("schedule_type", "all_periods")
    is_quincenal = pt.startswith("quincenal")

    if sched == "monthly_only":
        # Only deduct on mensual periods
        amt = monthly if not is_quincenal else 0.0
    elif sched == "biweekly_second_only":
        # Only on quincenal_2; nothing on quincenal_1 or mensual
        amt = monthly if pt == "quincenal_2" else 0.0
    else:
        # all_periods: split evenly for quincenal, full for mensual
        amt = monthly / 2 if is_quincenal else monthly

    return round(min(amt, remaining), 2)


async def loan_deduction_for_employee(employee_id: str, company_id: str, period_type: str | None) -> tuple[float, list[dict]]:
    """Sum the period's loan deduction across ALL active loans for the employee.
    Returns ``(total_deduction, active_loans)``.
    """
    active_loans = await db.loans.find(
        {"employee_id": employee_id, "company_id": company_id, "status": "active",
         "deduct_from_payroll": True, "remaining_balance": {"$gt": 0}},
        {"_id": 0}
    ).to_list(50)
    total = sum(loan_deduction_for_period(l, period_type) for l in active_loans)
    return round(total, 2), active_loans


async def register_loan_installment(loan_id: str, company_id: str, amount: float, period_id: str, payment_date: str) -> dict:
    """Register a payroll-driven installment payment on a loan.

    Decrements ``remaining_balance``, marks the next pending installment
    as paid in ``payment_schedule``, and auto-flips status to ``paid``
    when balance reaches 0.

    Idempotency note: callers should avoid double-calling this for the
    same (loan_id, period_id). The payroll ``pay_period`` flow only runs
    once per period when transitioning to paid status.
    """
    loan = await db.loans.find_one({"loan_id": loan_id, "company_id": company_id}, {"_id": 0})
    if not loan:
        return {"ok": False, "error": "loan_not_found"}
    if loan.get("status") != "active" or amount <= 0:
        return {"ok": False, "error": "loan_not_active_or_zero_amount"}

    new_total_paid = float(loan.get("total_paid", 0) or 0) + amount
    new_remaining = max(0.0, float(loan.get("remaining_balance", 0) or 0) - amount)

    schedule = loan.get("payment_schedule") or []
    for inst in schedule:
        if inst.get("status") == "pending":
            inst["status"] = "paid"
            inst["paid_date"] = payment_date
            inst["paid_amount"] = round(amount, 2)
            inst["period_id"] = period_id
            break

    payment = {
        "payment_id": f"pay_{uuid.uuid4().hex[:8]}",
        "amount": round(amount, 2),
        "payment_date": payment_date,
        "payment_type": "payroll",
        "period_id": period_id,
        "notes": f"Auto-deducción nómina {period_id}",
        "created_at": datetime.now(timezone.utc).isoformat(),
    }

    new_status = "paid" if new_remaining <= 0.01 else "active"
    await db.loans.update_one(
        {"loan_id": loan_id, "company_id": company_id},
        {
            "$set": {
                "total_paid": round(new_total_paid, 2),
                "remaining_balance": round(new_remaining, 2),
                "status": new_status,
                "payment_schedule": schedule,
                "updated_at": datetime.now(timezone.utc).isoformat(),
            },
            "$push": {"payments": payment},
        }
    )
    return {"ok": True, "remaining_balance": new_remaining, "status": new_status}


# ===================== EMPLOYEE-SPECIFIC LOAN ENDPOINT =====================

async def get_employee_loans_internal(employee_id: str, company_id: str):
    """Get all loans for a specific employee (internal function)"""
    loans = await db.loans.find(
        {"employee_id": employee_id, "company_id": company_id},
        {"_id": 0}
    ).sort("created_at", -1).to_list(100)
    
    # Calculate pending deduction for current period
    active_loans = [l for l in loans if l.get("status") == "active" and l.get("deduct_from_payroll")]
    pending_deduction = sum(l.get("monthly_payment", 0) for l in active_loans)
    
    return {
        "loans": loans,
        "pending_deduction": round(pending_deduction, 2),
        "total_balance": sum(l.get("remaining_balance", 0) for l in active_loans)
    }
