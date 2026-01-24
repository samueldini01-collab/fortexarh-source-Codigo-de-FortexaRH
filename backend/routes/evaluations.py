"""
Performance Evaluations - FortexaRH
Complete performance management with cycles, KPIs, 360° feedback, and improvement plans
"""
from fastapi import APIRouter, Request, HTTPException, Depends, Query
from fastapi.security import HTTPBearer
from pydantic import BaseModel
from typing import Optional, List, Dict, Any
from datetime import datetime, timezone, timedelta
import uuid
from io import BytesIO, StringIO
from fastapi.responses import StreamingResponse

router = APIRouter(prefix="/evaluations", tags=["Evaluations"])
security = HTTPBearer(auto_error=False)

db = None
_get_current_user_func = None


def init_router(database, auth_func):
    global db, _get_current_user_func
    db = database
    _get_current_user_func = auth_func


async def get_current_user(request: Request, credentials=Depends(security)):
    if _get_current_user_func is None:
        raise HTTPException(status_code=500, detail="Auth not initialized")
    return await _get_current_user_func(request, credentials)


# ==================== Scoring Scales ====================

NUMERIC_SCALE = [
    {"value": 1, "label": "1 - Insatisfactorio", "description": "No cumple con las expectativas mínimas"},
    {"value": 2, "label": "2 - Necesita Mejora", "description": "Cumple parcialmente, requiere desarrollo significativo"},
    {"value": 3, "label": "3 - Satisfactorio", "description": "Cumple con las expectativas del puesto"},
    {"value": 4, "label": "4 - Bueno", "description": "Supera las expectativas en varias áreas"},
    {"value": 5, "label": "5 - Excepcional", "description": "Desempeño sobresaliente en todas las áreas"},
]

DESCRIPTIVE_SCALE = [
    {"value": 1, "label": "Insatisfactorio", "color": "#ef4444"},
    {"value": 2, "label": "Necesita Mejora", "color": "#f97316"},
    {"value": 3, "label": "Satisfactorio", "color": "#eab308"},
    {"value": 4, "label": "Bueno", "color": "#22c55e"},
    {"value": 5, "label": "Excepcional", "color": "#10b981"},
]

DEFAULT_COMPETENCIES = [
    {"code": "performance", "name": "Desempeño Laboral", "description": "Calidad y cantidad de trabajo realizado", "weight": 25},
    {"code": "goals", "name": "Cumplimiento de Objetivos", "description": "Logro de metas establecidas", "weight": 25},
    {"code": "teamwork", "name": "Trabajo en Equipo", "description": "Colaboración y apoyo a compañeros", "weight": 15},
    {"code": "communication", "name": "Comunicación", "description": "Claridad y efectividad en la comunicación", "weight": 15},
    {"code": "initiative", "name": "Iniciativa", "description": "Proactividad y capacidad de innovar", "weight": 10},
    {"code": "punctuality", "name": "Puntualidad y Asistencia", "description": "Cumplimiento de horarios", "weight": 10},
]


# ==================== Models ====================

class EvaluationCycleCreate(BaseModel):
    name: str
    type: str = "annual"  # annual, semi_annual, quarterly, monthly
    start_date: str
    end_date: str
    self_evaluation_deadline: Optional[str] = None
    supervisor_evaluation_deadline: Optional[str] = None
    peer_evaluation_deadline: Optional[str] = None
    include_self_evaluation: bool = True
    include_peer_evaluation: bool = False
    competencies: Optional[List[Dict[str, Any]]] = None


class ObjectiveCreate(BaseModel):
    employee_id: str
    title: str
    description: Optional[str] = None
    target_value: Optional[float] = None
    target_unit: Optional[str] = None  # percentage, currency, quantity
    weight: float = 100  # Weight in overall evaluation
    due_date: Optional[str] = None
    cycle_id: Optional[str] = None


class ObjectiveUpdate(BaseModel):
    progress: Optional[float] = None  # 0-100
    actual_value: Optional[float] = None
    status: Optional[str] = None  # not_started, in_progress, completed, cancelled
    notes: Optional[str] = None


class EvaluationCreate(BaseModel):
    employee_id: str
    cycle_id: Optional[str] = None
    evaluation_type: str = "supervisor"  # supervisor, self, peer, 360
    period: str  # e.g., "Q1 2025", "2025"
    scores: List[Dict[str, Any]]  # [{competency: "performance", score: 4, comments: "..."}]
    overall_comments: Optional[str] = None
    strengths: Optional[List[str]] = None
    areas_for_improvement: Optional[List[str]] = None
    goals_for_next_period: Optional[List[str]] = None


class PeerEvaluationCreate(BaseModel):
    employee_id: str  # Employee being evaluated
    cycle_id: str
    scores: List[Dict[str, Any]]
    comments: Optional[str] = None
    relationship: str = "peer"  # peer, subordinate, cross_functional


class ImprovementPlanCreate(BaseModel):
    employee_id: str
    evaluation_id: Optional[str] = None
    title: str
    areas: List[str]  # Areas to improve
    actions: List[Dict[str, Any]]  # [{action: "...", deadline: "...", responsible: "..."}]
    start_date: str
    end_date: str
    follow_up_frequency: str = "monthly"  # weekly, bi_weekly, monthly


class ImprovementPlanUpdate(BaseModel):
    status: Optional[str] = None  # active, completed, cancelled
    progress_notes: Optional[str] = None
    action_updates: Optional[List[Dict[str, Any]]] = None


# ==================== Evaluation Cycles ====================

@router.get("/cycles")
async def get_evaluation_cycles(
    status: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    """Get all evaluation cycles"""
    query = {"company_id": current_user.get("company_id")}
    if status:
        query["status"] = status
    
    cycles = await db.evaluation_cycles.find(query, {"_id": 0}).sort("start_date", -1).to_list(100)
    return cycles


@router.get("/cycles/active")
async def get_active_cycle(current_user: dict = Depends(get_current_user)):
    """Get currently active evaluation cycle"""
    today = datetime.now().strftime("%Y-%m-%d")
    cycle = await db.evaluation_cycles.find_one({
        "company_id": current_user.get("company_id"),
        "start_date": {"$lte": today},
        "end_date": {"$gte": today},
        "status": "active"
    }, {"_id": 0})
    return cycle


@router.post("/cycles")
async def create_evaluation_cycle(data: EvaluationCycleCreate, current_user: dict = Depends(get_current_user)):
    """Create a new evaluation cycle"""
    company_id = current_user.get("company_id")
    
    # Check for overlapping cycles
    overlap = await db.evaluation_cycles.find_one({
        "company_id": company_id,
        "status": "active",
        "$or": [
            {"start_date": {"$lte": data.end_date}, "end_date": {"$gte": data.start_date}}
        ]
    })
    if overlap:
        raise HTTPException(status_code=400, detail="Ya existe un ciclo activo en estas fechas")
    
    cycle_id = f"cycle_{uuid.uuid4().hex[:12]}"
    cycle = {
        "cycle_id": cycle_id,
        "company_id": company_id,
        "name": data.name,
        "type": data.type,
        "start_date": data.start_date,
        "end_date": data.end_date,
        "self_evaluation_deadline": data.self_evaluation_deadline,
        "supervisor_evaluation_deadline": data.supervisor_evaluation_deadline,
        "peer_evaluation_deadline": data.peer_evaluation_deadline,
        "include_self_evaluation": data.include_self_evaluation,
        "include_peer_evaluation": data.include_peer_evaluation,
        "competencies": data.competencies or DEFAULT_COMPETENCIES,
        "status": "active",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "created_by": current_user["user_id"]
    }
    await db.evaluation_cycles.insert_one(cycle)
    
    return {"cycle_id": cycle_id, "message": "Ciclo de evaluación creado"}


@router.put("/cycles/{cycle_id}")
async def update_evaluation_cycle(
    cycle_id: str,
    data: EvaluationCycleCreate,
    current_user: dict = Depends(get_current_user)
):
    """Update an evaluation cycle"""
    result = await db.evaluation_cycles.update_one(
        {"cycle_id": cycle_id, "company_id": current_user.get("company_id")},
        {"$set": {
            "name": data.name,
            "type": data.type,
            "start_date": data.start_date,
            "end_date": data.end_date,
            "self_evaluation_deadline": data.self_evaluation_deadline,
            "supervisor_evaluation_deadline": data.supervisor_evaluation_deadline,
            "peer_evaluation_deadline": data.peer_evaluation_deadline,
            "include_self_evaluation": data.include_self_evaluation,
            "include_peer_evaluation": data.include_peer_evaluation,
            "competencies": data.competencies,
            "updated_at": datetime.now(timezone.utc).isoformat()
        }}
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Ciclo no encontrado")
    return {"message": "Ciclo actualizado"}


@router.put("/cycles/{cycle_id}/close")
async def close_evaluation_cycle(cycle_id: str, current_user: dict = Depends(get_current_user)):
    """Close an evaluation cycle"""
    result = await db.evaluation_cycles.update_one(
        {"cycle_id": cycle_id, "company_id": current_user.get("company_id")},
        {"$set": {"status": "closed", "closed_at": datetime.now(timezone.utc).isoformat()}}
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Ciclo no encontrado")
    return {"message": "Ciclo cerrado"}


# ==================== Objectives / KPIs ====================

@router.get("/objectives")
async def get_objectives(
    employee_id: Optional[str] = None,
    cycle_id: Optional[str] = None,
    status: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    """Get objectives/KPIs"""
    query = {"company_id": current_user.get("company_id")}
    if employee_id:
        query["employee_id"] = employee_id
    if cycle_id:
        query["cycle_id"] = cycle_id
    if status:
        query["status"] = status
    
    objectives = await db.objectives.find(query, {"_id": 0}).to_list(500)
    return objectives


@router.get("/objectives/{employee_id}")
async def get_employee_objectives(
    employee_id: str,
    current_user: dict = Depends(get_current_user)
):
    """Get all objectives for an employee"""
    objectives = await db.objectives.find({
        "company_id": current_user.get("company_id"),
        "employee_id": employee_id
    }, {"_id": 0}).sort("created_at", -1).to_list(100)
    
    # Calculate summary
    total = len(objectives)
    completed = sum(1 for o in objectives if o.get("status") == "completed")
    in_progress = sum(1 for o in objectives if o.get("status") == "in_progress")
    avg_progress = sum(o.get("progress", 0) for o in objectives) / max(total, 1)
    
    return {
        "employee_id": employee_id,
        "objectives": objectives,
        "summary": {
            "total": total,
            "completed": completed,
            "in_progress": in_progress,
            "not_started": total - completed - in_progress,
            "average_progress": round(avg_progress, 1)
        }
    }


@router.post("/objectives")
async def create_objective(data: ObjectiveCreate, current_user: dict = Depends(get_current_user)):
    """Create a new objective/KPI for an employee"""
    company_id = current_user.get("company_id")
    
    employee = await db.employees.find_one(
        {"employee_id": data.employee_id, "company_id": company_id},
        {"_id": 0}
    )
    if not employee:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")
    
    objective_id = f"obj_{uuid.uuid4().hex[:12]}"
    objective = {
        "objective_id": objective_id,
        "company_id": company_id,
        "employee_id": data.employee_id,
        "employee_name": f"{employee['first_name']} {employee['last_name']}",
        "department": employee.get("department"),
        "title": data.title,
        "description": data.description,
        "target_value": data.target_value,
        "target_unit": data.target_unit,
        "actual_value": None,
        "weight": data.weight,
        "due_date": data.due_date,
        "cycle_id": data.cycle_id,
        "progress": 0,
        "status": "not_started",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "created_by": current_user["user_id"]
    }
    await db.objectives.insert_one(objective)
    
    return {"objective_id": objective_id, "message": "Objetivo creado"}


@router.put("/objectives/{objective_id}")
async def update_objective(
    objective_id: str,
    data: ObjectiveUpdate,
    current_user: dict = Depends(get_current_user)
):
    """Update objective progress"""
    update_data = {"updated_at": datetime.now(timezone.utc).isoformat()}
    
    if data.progress is not None:
        update_data["progress"] = data.progress
        if data.progress >= 100:
            update_data["status"] = "completed"
            update_data["completed_at"] = datetime.now(timezone.utc).isoformat()
        elif data.progress > 0:
            update_data["status"] = "in_progress"
    
    if data.actual_value is not None:
        update_data["actual_value"] = data.actual_value
    if data.status:
        update_data["status"] = data.status
    if data.notes:
        update_data["notes"] = data.notes
    
    result = await db.objectives.update_one(
        {"objective_id": objective_id, "company_id": current_user.get("company_id")},
        {"$set": update_data}
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Objetivo no encontrado")
    
    return {"message": "Objetivo actualizado"}


@router.delete("/objectives/{objective_id}")
async def delete_objective(objective_id: str, current_user: dict = Depends(get_current_user)):
    """Delete an objective"""
    result = await db.objectives.delete_one({
        "objective_id": objective_id,
        "company_id": current_user.get("company_id")
    })
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Objetivo no encontrado")
    return {"message": "Objetivo eliminado"}


# ==================== Evaluations ====================

@router.get("")
async def get_evaluations(
    employee_id: Optional[str] = None,
    cycle_id: Optional[str] = None,
    evaluation_type: Optional[str] = None,
    status: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    """Get evaluations with filters"""
    query = {"company_id": current_user.get("company_id")}
    if employee_id:
        query["employee_id"] = employee_id
    if cycle_id:
        query["cycle_id"] = cycle_id
    if evaluation_type:
        query["evaluation_type"] = evaluation_type
    if status:
        query["status"] = status
    
    evaluations = await db.evaluations.find(query, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return evaluations


@router.get("/employee/{employee_id}")
async def get_employee_evaluations(
    employee_id: str,
    current_user: dict = Depends(get_current_user)
):
    """Get all evaluations for an employee with history"""
    company_id = current_user.get("company_id")
    
    employee = await db.employees.find_one(
        {"employee_id": employee_id, "company_id": company_id},
        {"_id": 0}
    )
    if not employee:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")
    
    evaluations = await db.evaluations.find({
        "company_id": company_id,
        "employee_id": employee_id
    }, {"_id": 0}).sort("created_at", -1).to_list(50)
    
    # Calculate trends
    scores_over_time = []
    for eval in evaluations:
        if eval.get("overall_score"):
            scores_over_time.append({
                "period": eval.get("period"),
                "score": eval["overall_score"],
                "type": eval.get("evaluation_type")
            })
    
    return {
        "employee": {
            "employee_id": employee_id,
            "name": f"{employee['first_name']} {employee['last_name']}",
            "department": employee.get("department"),
            "position": employee.get("position")
        },
        "evaluations": evaluations,
        "score_history": scores_over_time[:10],
        "average_score": sum(s["score"] for s in scores_over_time) / max(len(scores_over_time), 1) if scores_over_time else None
    }


@router.post("")
async def create_evaluation(data: EvaluationCreate, current_user: dict = Depends(get_current_user)):
    """Create a new evaluation"""
    company_id = current_user.get("company_id")
    
    employee = await db.employees.find_one(
        {"employee_id": data.employee_id, "company_id": company_id},
        {"_id": 0}
    )
    if not employee:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")
    
    # Calculate overall score
    total_score = 0
    total_weight = 0
    for score in data.scores:
        weight = score.get("weight", 1)
        total_score += score.get("score", 0) * weight
        total_weight += weight
    
    overall_score = round(total_score / max(total_weight, 1), 2)
    
    # Determine descriptive rating
    if overall_score >= 4.5:
        rating = "Excepcional"
    elif overall_score >= 3.5:
        rating = "Bueno"
    elif overall_score >= 2.5:
        rating = "Satisfactorio"
    elif overall_score >= 1.5:
        rating = "Necesita Mejora"
    else:
        rating = "Insatisfactorio"
    
    evaluation_id = f"eval_{uuid.uuid4().hex[:12]}"
    evaluation = {
        "evaluation_id": evaluation_id,
        "company_id": company_id,
        "employee_id": data.employee_id,
        "employee_name": f"{employee['first_name']} {employee['last_name']}",
        "department": employee.get("department"),
        "position": employee.get("position"),
        "cycle_id": data.cycle_id,
        "evaluation_type": data.evaluation_type,
        "period": data.period,
        "scores": data.scores,
        "overall_score": overall_score,
        "rating": rating,
        "overall_comments": data.overall_comments,
        "strengths": data.strengths or [],
        "areas_for_improvement": data.areas_for_improvement or [],
        "goals_for_next_period": data.goals_for_next_period or [],
        "status": "draft",
        "evaluator_id": current_user["user_id"],
        "evaluator_name": current_user.get("name"),
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.evaluations.insert_one(evaluation)
    
    return {
        "evaluation_id": evaluation_id,
        "overall_score": overall_score,
        "rating": rating,
        "message": "Evaluación creada"
    }


@router.put("/{evaluation_id}")
async def update_evaluation(
    evaluation_id: str,
    data: EvaluationCreate,
    current_user: dict = Depends(get_current_user)
):
    """Update an evaluation"""
    company_id = current_user.get("company_id")
    
    existing = await db.evaluations.find_one({
        "evaluation_id": evaluation_id,
        "company_id": company_id
    })
    if not existing:
        raise HTTPException(status_code=404, detail="Evaluación no encontrada")
    
    if existing.get("status") == "finalized":
        raise HTTPException(status_code=400, detail="No se puede modificar una evaluación finalizada")
    
    # Recalculate score
    total_score = 0
    total_weight = 0
    for score in data.scores:
        weight = score.get("weight", 1)
        total_score += score.get("score", 0) * weight
        total_weight += weight
    
    overall_score = round(total_score / max(total_weight, 1), 2)
    
    if overall_score >= 4.5:
        rating = "Excepcional"
    elif overall_score >= 3.5:
        rating = "Bueno"
    elif overall_score >= 2.5:
        rating = "Satisfactorio"
    elif overall_score >= 1.5:
        rating = "Necesita Mejora"
    else:
        rating = "Insatisfactorio"
    
    await db.evaluations.update_one(
        {"evaluation_id": evaluation_id},
        {"$set": {
            "scores": data.scores,
            "overall_score": overall_score,
            "rating": rating,
            "overall_comments": data.overall_comments,
            "strengths": data.strengths,
            "areas_for_improvement": data.areas_for_improvement,
            "goals_for_next_period": data.goals_for_next_period,
            "updated_at": datetime.now(timezone.utc).isoformat()
        }}
    )
    
    return {"overall_score": overall_score, "rating": rating, "message": "Evaluación actualizada"}


@router.put("/{evaluation_id}/finalize")
async def finalize_evaluation(evaluation_id: str, current_user: dict = Depends(get_current_user)):
    """Finalize an evaluation (lock it)"""
    result = await db.evaluations.update_one(
        {"evaluation_id": evaluation_id, "company_id": current_user.get("company_id")},
        {"$set": {
            "status": "finalized",
            "finalized_at": datetime.now(timezone.utc).isoformat(),
            "finalized_by": current_user["user_id"]
        }}
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Evaluación no encontrada")
    return {"message": "Evaluación finalizada"}


# ==================== 360° Feedback ====================

@router.post("/360/request")
async def request_360_feedback(
    employee_id: str = Query(...),
    cycle_id: str = Query(...),
    evaluator_ids: List[str] = Query(...),
    current_user: dict = Depends(get_current_user)
):
    """Request 360° feedback from peers/subordinates"""
    company_id = current_user.get("company_id")
    
    employee = await db.employees.find_one(
        {"employee_id": employee_id, "company_id": company_id},
        {"_id": 0}
    )
    if not employee:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")
    
    requests_created = []
    for evaluator_id in evaluator_ids:
        request_id = f"360req_{uuid.uuid4().hex[:12]}"
        await db.feedback_requests.insert_one({
            "request_id": request_id,
            "company_id": company_id,
            "employee_id": employee_id,
            "employee_name": f"{employee['first_name']} {employee['last_name']}",
            "evaluator_id": evaluator_id,
            "cycle_id": cycle_id,
            "status": "pending",
            "created_at": datetime.now(timezone.utc).isoformat()
        })
        requests_created.append(request_id)
        
        # Create notification
        await db.notifications.insert_one({
            "notification_id": f"notif_{uuid.uuid4().hex[:12]}",
            "company_id": company_id,
            "user_id": evaluator_id,
            "type": "feedback_request",
            "title": "Solicitud de Evaluación 360°",
            "message": f"Se te ha solicitado evaluar a {employee['first_name']} {employee['last_name']}",
            "reference_id": request_id,
            "read": False,
            "created_at": datetime.now(timezone.utc).isoformat()
        })
    
    return {"requests_created": len(requests_created), "message": "Solicitudes enviadas"}


@router.get("/360/pending")
async def get_pending_360_requests(current_user: dict = Depends(get_current_user)):
    """Get pending 360° feedback requests for current user"""
    requests = await db.feedback_requests.find({
        "company_id": current_user.get("company_id"),
        "evaluator_id": current_user["user_id"],
        "status": "pending"
    }, {"_id": 0}).to_list(50)
    return requests


@router.post("/360/submit")
async def submit_360_feedback(data: PeerEvaluationCreate, current_user: dict = Depends(get_current_user)):
    """Submit 360° feedback"""
    company_id = current_user.get("company_id")
    
    # Calculate score
    total_score = sum(s.get("score", 0) for s in data.scores) / max(len(data.scores), 1)
    
    feedback_id = f"360fb_{uuid.uuid4().hex[:12]}"
    feedback = {
        "feedback_id": feedback_id,
        "company_id": company_id,
        "employee_id": data.employee_id,
        "evaluator_id": current_user["user_id"],
        "evaluator_name": current_user.get("name"),
        "cycle_id": data.cycle_id,
        "relationship": data.relationship,
        "scores": data.scores,
        "overall_score": round(total_score, 2),
        "comments": data.comments,
        "is_anonymous": True,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.peer_evaluations.insert_one(feedback)
    
    # Update request status
    await db.feedback_requests.update_one(
        {
            "company_id": company_id,
            "employee_id": data.employee_id,
            "evaluator_id": current_user["user_id"],
            "cycle_id": data.cycle_id
        },
        {"$set": {"status": "completed", "completed_at": datetime.now(timezone.utc).isoformat()}}
    )
    
    return {"feedback_id": feedback_id, "message": "Evaluación 360° enviada"}


@router.get("/360/summary/{employee_id}")
async def get_360_summary(
    employee_id: str,
    cycle_id: str = Query(...),
    current_user: dict = Depends(get_current_user)
):
    """Get 360° feedback summary (anonymized)"""
    company_id = current_user.get("company_id")
    
    feedbacks = await db.peer_evaluations.find({
        "company_id": company_id,
        "employee_id": employee_id,
        "cycle_id": cycle_id
    }, {"_id": 0, "evaluator_id": 0, "evaluator_name": 0}).to_list(50)
    
    if not feedbacks:
        return {"message": "No hay evaluaciones 360° para este ciclo"}
    
    # Aggregate scores by competency
    competency_scores = {}
    for fb in feedbacks:
        for score in fb.get("scores", []):
            comp = score.get("competency")
            if comp not in competency_scores:
                competency_scores[comp] = []
            competency_scores[comp].append(score.get("score", 0))
    
    # Calculate averages
    competency_averages = {
        comp: round(sum(scores) / len(scores), 2)
        for comp, scores in competency_scores.items()
    }
    
    overall_avg = sum(fb["overall_score"] for fb in feedbacks) / len(feedbacks)
    
    return {
        "employee_id": employee_id,
        "cycle_id": cycle_id,
        "total_evaluators": len(feedbacks),
        "overall_average": round(overall_avg, 2),
        "competency_averages": competency_averages,
        "anonymous_comments": [fb.get("comments") for fb in feedbacks if fb.get("comments")]
    }


# ==================== Improvement Plans ====================

@router.get("/improvement-plans")
async def get_improvement_plans(
    employee_id: Optional[str] = None,
    status: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    """Get improvement plans"""
    query = {"company_id": current_user.get("company_id")}
    if employee_id:
        query["employee_id"] = employee_id
    if status:
        query["status"] = status
    
    plans = await db.improvement_plans.find(query, {"_id": 0}).sort("created_at", -1).to_list(200)
    return plans


@router.post("/improvement-plans")
async def create_improvement_plan(data: ImprovementPlanCreate, current_user: dict = Depends(get_current_user)):
    """Create an improvement plan"""
    company_id = current_user.get("company_id")
    
    employee = await db.employees.find_one(
        {"employee_id": data.employee_id, "company_id": company_id},
        {"_id": 0}
    )
    if not employee:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")
    
    plan_id = f"plan_{uuid.uuid4().hex[:12]}"
    plan = {
        "plan_id": plan_id,
        "company_id": company_id,
        "employee_id": data.employee_id,
        "employee_name": f"{employee['first_name']} {employee['last_name']}",
        "department": employee.get("department"),
        "evaluation_id": data.evaluation_id,
        "title": data.title,
        "areas": data.areas,
        "actions": data.actions,
        "start_date": data.start_date,
        "end_date": data.end_date,
        "follow_up_frequency": data.follow_up_frequency,
        "status": "active",
        "progress_history": [],
        "created_at": datetime.now(timezone.utc).isoformat(),
        "created_by": current_user["user_id"]
    }
    await db.improvement_plans.insert_one(plan)
    
    return {"plan_id": plan_id, "message": "Plan de mejora creado"}


@router.put("/improvement-plans/{plan_id}")
async def update_improvement_plan(
    plan_id: str,
    data: ImprovementPlanUpdate,
    current_user: dict = Depends(get_current_user)
):
    """Update improvement plan progress"""
    update_data = {"updated_at": datetime.now(timezone.utc).isoformat()}
    
    if data.status:
        update_data["status"] = data.status
        if data.status == "completed":
            update_data["completed_at"] = datetime.now(timezone.utc).isoformat()
    
    if data.action_updates:
        update_data["actions"] = data.action_updates
    
    # Add progress note to history
    push_data = {}
    if data.progress_notes:
        push_data["progress_history"] = {
            "note": data.progress_notes,
            "added_by": current_user["user_id"],
            "added_at": datetime.now(timezone.utc).isoformat()
        }
    
    update_ops = {"$set": update_data}
    if push_data:
        update_ops["$push"] = push_data
    
    result = await db.improvement_plans.update_one(
        {"plan_id": plan_id, "company_id": current_user.get("company_id")},
        update_ops
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Plan no encontrado")
    
    return {"message": "Plan actualizado"}


# ==================== Reports & Analytics ====================

@router.get("/analytics/dashboard")
async def get_evaluations_dashboard(
    cycle_id: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    """Get evaluation analytics dashboard"""
    company_id = current_user.get("company_id")
    
    match_query = {"company_id": company_id}
    if cycle_id:
        match_query["cycle_id"] = cycle_id
    
    # Score distribution
    distribution_pipeline = [
        {"$match": match_query},
        {
            "$bucket": {
                "groupBy": "$overall_score",
                "boundaries": [0, 2, 3, 4, 5, 6],
                "default": "Other",
                "output": {"count": {"$sum": 1}}
            }
        }
    ]
    distribution = await db.evaluations.aggregate(distribution_pipeline).to_list(10)
    
    # Department averages
    dept_pipeline = [
        {"$match": match_query},
        {
            "$group": {
                "_id": "$department",
                "average_score": {"$avg": "$overall_score"},
                "count": {"$sum": 1}
            }
        },
        {"$sort": {"average_score": -1}}
    ]
    dept_stats = await db.evaluations.aggregate(dept_pipeline).to_list(50)
    
    # Top performers
    top_pipeline = [
        {"$match": match_query},
        {"$sort": {"overall_score": -1}},
        {"$limit": 10},
        {
            "$project": {
                "_id": 0,
                "employee_id": 1,
                "employee_name": 1,
                "department": 1,
                "overall_score": 1,
                "rating": 1
            }
        }
    ]
    top_performers = await db.evaluations.aggregate(top_pipeline).to_list(10)
    
    # Needs improvement
    needs_improvement = await db.evaluations.find({
        **match_query,
        "overall_score": {"$lt": 2.5}
    }, {"_id": 0, "employee_id": 1, "employee_name": 1, "department": 1, "overall_score": 1}).to_list(20)
    
    # Overall stats
    total_evals = await db.evaluations.count_documents(match_query)
    avg_score_result = await db.evaluations.aggregate([
        {"$match": match_query},
        {"$group": {"_id": None, "avg": {"$avg": "$overall_score"}}}
    ]).to_list(1)
    avg_score = avg_score_result[0]["avg"] if avg_score_result else 0
    
    return {
        "total_evaluations": total_evals,
        "average_score": round(avg_score, 2) if avg_score else 0,
        "score_distribution": distribution,
        "department_stats": dept_stats,
        "top_performers": top_performers,
        "needs_improvement": needs_improvement
    }


@router.get("/scales")
async def get_evaluation_scales():
    """Get available evaluation scales"""
    return {
        "numeric": NUMERIC_SCALE,
        "descriptive": DESCRIPTIVE_SCALE
    }


@router.get("/competencies")
async def get_competencies(current_user: dict = Depends(get_current_user)):
    """Get company competencies or defaults"""
    custom = await db.competencies.find(
        {"company_id": current_user.get("company_id")},
        {"_id": 0}
    ).to_list(50)
    
    return custom if custom else DEFAULT_COMPETENCIES


@router.get("/export")
async def export_evaluations(
    cycle_id: Optional[str] = None,
    format: str = Query("csv"),
    current_user: dict = Depends(get_current_user)
):
    """Export evaluations to CSV or Excel"""
    query = {"company_id": current_user.get("company_id")}
    if cycle_id:
        query["cycle_id"] = cycle_id
    
    evaluations = await db.evaluations.find(query, {"_id": 0}).to_list(10000)
    
    if format == "csv":
        import csv
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "Empleado", "Departamento", "Posición", "Período", "Tipo",
            "Puntuación", "Calificación", "Fortalezas", "Áreas de Mejora"
        ])
        
        for e in evaluations:
            writer.writerow([
                e.get("employee_name", ""),
                e.get("department", ""),
                e.get("position", ""),
                e.get("period", ""),
                e.get("evaluation_type", ""),
                e.get("overall_score", ""),
                e.get("rating", ""),
                "; ".join(e.get("strengths", [])),
                "; ".join(e.get("areas_for_improvement", []))
            ])
        
        output.seek(0)
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=evaluaciones.csv"}
        )
    
    elif format == "excel":
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Evaluaciones"
        
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        
        headers = ["Empleado", "Departamento", "Posición", "Período", "Tipo",
                   "Puntuación", "Calificación", "Fortalezas", "Áreas de Mejora"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
        
        for row, e in enumerate(evaluations, 2):
            ws.cell(row=row, column=1, value=e.get("employee_name", ""))
            ws.cell(row=row, column=2, value=e.get("department", ""))
            ws.cell(row=row, column=3, value=e.get("position", ""))
            ws.cell(row=row, column=4, value=e.get("period", ""))
            ws.cell(row=row, column=5, value=e.get("evaluation_type", ""))
            ws.cell(row=row, column=6, value=e.get("overall_score", ""))
            ws.cell(row=row, column=7, value=e.get("rating", ""))
            ws.cell(row=row, column=8, value="; ".join(e.get("strengths", [])))
            ws.cell(row=row, column=9, value="; ".join(e.get("areas_for_improvement", [])))
        
        for col in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 40)
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=evaluaciones.xlsx"}
        )
    
    raise HTTPException(status_code=400, detail="Formato no soportado")
