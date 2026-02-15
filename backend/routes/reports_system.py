"""
Advanced Reports System - FortexaRH
Endpoints for the reporting system. Definitions and generators are in services/.
"""
from fastapi import APIRouter, Depends, HTTPException, Request, Query
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPBearer
from typing import Callable, Optional
from datetime import datetime, timezone
import io
import uuid

# PDF Generation
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER

from services.report_catalog import REPORT_CATEGORIES, REPORT_DEFINITIONS
from services.report_generators import generate_report_data, format_currency

router = APIRouter(prefix="/reports-system", tags=["Reports System"])
security = HTTPBearer(auto_error=False)

db = None
_get_current_user_func: Callable = None


def init_router(database, auth_dependency: Callable):
    global db, _get_current_user_func
    db = database
    _get_current_user_func = auth_dependency


async def get_current_user(request: Request, credentials=Depends(security)):
    if _get_current_user_func is None:
        raise HTTPException(status_code=500, detail="Auth not initialized")
    return await _get_current_user_func(request, credentials)


# ============== MODELS ==============
from models.system import ReportFilter, ReportRequest, SavedReportConfig


# ============== ENDPOINTS ==============

@router.get("/catalog")
async def get_report_catalog(current_user: dict = Depends(get_current_user)):
    """Get full catalog of available reports organized by category"""
    catalog = {}
    for cat_id, cat_info in REPORT_CATEGORIES.items():
        catalog[cat_id] = {**cat_info, "reports": []}
    for report_id, report in REPORT_DEFINITIONS.items():
        category = report["category"]
        if category in catalog:
            catalog[category]["reports"].append({
                "id": report["id"],
                "name": report["name"],
                "description": report["description"],
                "filters": report["filters"],
                "supports_preview": report.get("supports_preview", True),
                "supports_pdf": report.get("supports_pdf", True),
                "supports_excel": report.get("supports_excel", True)
            })
    return {"categories": catalog, "total_reports": len(REPORT_DEFINITIONS)}


@router.get("/definition/{report_id}")
async def get_report_definition(report_id: str, current_user: dict = Depends(get_current_user)):
    if report_id not in REPORT_DEFINITIONS:
        raise HTTPException(status_code=404, detail="Reporte no encontrado")
    return REPORT_DEFINITIONS[report_id]


@router.post("/preview")
async def preview_report(request: ReportRequest, current_user: dict = Depends(get_current_user)):
    if request.report_id not in REPORT_DEFINITIONS:
        raise HTTPException(status_code=404, detail="Reporte no encontrado")
    company_id = current_user.get("company_id")
    report_def = REPORT_DEFINITIONS[request.report_id]
    filters = {f.get("field"): f.get("value") for f in request.filters}
    data = await generate_report_data(
        request.report_id, company_id, filters,
        request.page, request.page_size, request.sort_by, request.sort_order
    )
    return {
        "report_id": request.report_id,
        "report_name": report_def["name"],
        "filters_applied": filters,
        "columns": request.columns or report_def["columns"],
        "data": data["rows"],
        "totals": data.get("totals", {}),
        "summary": data.get("summary", {}),
        "pagination": {
            "page": request.page,
            "page_size": request.page_size,
            "total_rows": data.get("total_count", len(data["rows"])),
            "total_pages": (data.get("total_count", len(data["rows"])) + request.page_size - 1) // request.page_size
        },
        "generated_at": datetime.now(timezone.utc).isoformat()
    }


@router.post("/export/pdf")
async def export_report_pdf(request: ReportRequest, current_user: dict = Depends(get_current_user)):
    if request.report_id not in REPORT_DEFINITIONS:
        raise HTTPException(status_code=404, detail="Reporte no encontrado")
    company_id = current_user.get("company_id")
    report_def = REPORT_DEFINITIONS[request.report_id]
    company = await db.companies.find_one({"company_id": company_id}, {"_id": 0, "name": 1, "rnc": 1})
    company_name = company.get("name", "FortexaRH") if company else "FortexaRH"
    filters = {f.get("field"): f.get("value") for f in request.filters}
    data = await generate_report_data(request.report_id, company_id, filters, 1, 10000, request.sort_by, request.sort_order)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=16, alignment=TA_CENTER, spaceAfter=12)
    subtitle_style = ParagraphStyle('CustomSubtitle', parent=styles['Normal'], fontSize=10, alignment=TA_CENTER, textColor=colors.gray)
    elements = []
    elements.append(Paragraph(f"<b>{company_name}</b>", title_style))
    elements.append(Paragraph(report_def["name"], styles['Heading2']))
    elements.append(Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style))
    elements.append(Spacer(1, 0.3*inch))

    if filters:
        filter_text = " | ".join([f"{k}: {v}" for k, v in filters.items() if v])
        if filter_text:
            elements.append(Paragraph(f"<b>Filtros:</b> {filter_text}", styles['Normal']))
            elements.append(Spacer(1, 0.2*inch))

    columns = request.columns or report_def["columns"]
    header_row = [col.replace("_", " ").title() for col in columns]
    table_data = [header_row]
    for row in data["rows"]:
        table_row = []
        for col in columns:
            value = row.get(col, "")
            if isinstance(value, float):
                value = format_currency(value) if any(k in col for k in ("salary", "amount", "cost", "pay")) else f"{value:.2f}"
            elif isinstance(value, bool):
                value = "Sí" if value else "No"
            table_row.append(str(value) if value is not None else "")
        table_data.append(table_row)

    col_widths = [1.2*inch] * len(columns)
    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('BACKGROUND', (0, 1), (-1, -1), colors.white),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
    ]))
    elements.append(table)

    if data.get("totals"):
        elements.append(Spacer(1, 0.2*inch))
        elements.append(Paragraph("<b>Totales</b>", styles['Heading3']))
        totals_text = " | ".join([f"{k}: {format_currency(v) if isinstance(v, (int, float)) else v}" for k, v in data["totals"].items()])
        elements.append(Paragraph(totals_text, styles['Normal']))

    elements.append(Spacer(1, 0.3*inch))
    elements.append(Paragraph(f"Total de registros: {len(data['rows'])}", subtitle_style))
    elements.append(Paragraph(f"Generado por: {current_user.get('name', 'Usuario')}", subtitle_style))
    doc.build(elements)
    buffer.seek(0)
    await record_report_generation(company_id, current_user, request.report_id, filters, "pdf", len(data["rows"]))
    filename = f"{report_def['name'].replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={filename}"})


@router.post("/export/excel")
async def export_report_excel(request: ReportRequest, current_user: dict = Depends(get_current_user)):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl no está instalado")

    if request.report_id not in REPORT_DEFINITIONS:
        raise HTTPException(status_code=404, detail="Reporte no encontrado")
    company_id = current_user.get("company_id")
    report_def = REPORT_DEFINITIONS[request.report_id]
    filters = {f.get("field"): f.get("value") for f in request.filters}
    data = await generate_report_data(request.report_id, company_id, filters, 1, 10000, request.sort_by, request.sort_order)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = report_def["name"][:31]
    header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    ws['A1'] = report_def["name"]
    ws['A1'].font = Font(size=14, bold=True)
    ws['A2'] = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    columns = request.columns or report_def["columns"]
    for col_idx, col in enumerate(columns, 1):
        cell = ws.cell(row=4, column=col_idx, value=col.replace("_", " ").title())
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = border
    for row_idx, row in enumerate(data["rows"], 5):
        for col_idx, col in enumerate(columns, 1):
            value = row.get(col, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = border
            if isinstance(value, (int, float)):
                cell.alignment = Alignment(horizontal='right')
    for col_idx, col in enumerate(columns, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 15

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    await record_report_generation(company_id, current_user, request.report_id, filters, "excel", len(data["rows"]))
    filename = f"{report_def['name'].replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename={filename}"})


@router.post("/export/csv")
async def export_report_csv(request: ReportRequest, current_user: dict = Depends(get_current_user)):
    import csv
    if request.report_id not in REPORT_DEFINITIONS:
        raise HTTPException(status_code=404, detail="Reporte no encontrado")
    company_id = current_user.get("company_id")
    report_def = REPORT_DEFINITIONS[request.report_id]
    filters = {f.get("field"): f.get("value") for f in request.filters}
    data = await generate_report_data(request.report_id, company_id, filters, 1, 10000, request.sort_by, request.sort_order)
    columns = request.columns or report_def["columns"]
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([col.replace("_", " ").title() for col in columns])
    for row in data["rows"]:
        writer.writerow([row.get(col, "") for col in columns])
    await record_report_generation(company_id, current_user, request.report_id, filters, "csv", len(data["rows"]))
    filename = f"{report_def['name'].replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
    output = io.BytesIO(buffer.getvalue().encode('utf-8-sig'))
    return StreamingResponse(output, media_type="text/csv", headers={"Content-Disposition": f"attachment; filename={filename}"})


# ============== SAVED REPORTS ==============

@router.post("/saved")
async def save_report_config(config: SavedReportConfig, current_user: dict = Depends(get_current_user)):
    company_id = current_user.get("company_id")
    user_id = current_user.get("user_id")
    saved_report = {
        "saved_report_id": f"sr_{uuid.uuid4().hex[:12]}",
        "company_id": company_id,
        "user_id": user_id,
        "name": config.name,
        "description": config.description,
        "report_id": config.report_id,
        "filters": config.filters,
        "columns": config.columns,
        "sort_by": config.sort_by,
        "sort_order": config.sort_order,
        "is_favorite": config.is_favorite,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "updated_at": datetime.now(timezone.utc).isoformat()
    }
    await db.saved_reports.insert_one(saved_report)
    return {"message": "Configuración guardada", "saved_report_id": saved_report["saved_report_id"]}


@router.get("/saved")
async def get_saved_reports(current_user: dict = Depends(get_current_user)):
    company_id = current_user.get("company_id")
    user_id = current_user.get("user_id")
    saved = await db.saved_reports.find(
        {"company_id": company_id, "user_id": user_id}, {"_id": 0}
    ).sort("created_at", -1).to_list(100)
    for s in saved:
        if s["report_id"] in REPORT_DEFINITIONS:
            s["report_name"] = REPORT_DEFINITIONS[s["report_id"]]["name"]
            s["report_category"] = REPORT_DEFINITIONS[s["report_id"]]["category"]
    return saved


@router.get("/saved/{saved_report_id}")
async def get_saved_report(saved_report_id: str, current_user: dict = Depends(get_current_user)):
    company_id = current_user.get("company_id")
    saved = await db.saved_reports.find_one(
        {"saved_report_id": saved_report_id, "company_id": company_id}, {"_id": 0}
    )
    if not saved:
        raise HTTPException(status_code=404, detail="Configuración no encontrada")
    return saved


@router.delete("/saved/{saved_report_id}")
async def delete_saved_report(saved_report_id: str, current_user: dict = Depends(get_current_user)):
    company_id = current_user.get("company_id")
    result = await db.saved_reports.delete_one({"saved_report_id": saved_report_id, "company_id": company_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Configuración no encontrada")
    return {"message": "Configuración eliminada"}


# ============== TRACEABILITY ==============

@router.get("/history")
async def get_report_history(limit: int = Query(50, le=200), current_user: dict = Depends(get_current_user)):
    company_id = current_user.get("company_id")
    history = await db.report_history.find(
        {"company_id": company_id}, {"_id": 0}
    ).sort("generated_at", -1).limit(limit).to_list(limit)
    for h in history:
        if h.get("report_id") in REPORT_DEFINITIONS:
            h["report_name"] = REPORT_DEFINITIONS[h["report_id"]]["name"]
    return history


async def record_report_generation(company_id: str, user: dict, report_id: str, filters: dict, format: str, row_count: int):
    record = {
        "history_id": f"rh_{uuid.uuid4().hex[:12]}",
        "company_id": company_id,
        "user_id": user.get("user_id"),
        "user_name": user.get("name", "Usuario"),
        "report_id": report_id,
        "filters_applied": filters,
        "export_format": format,
        "row_count": row_count,
        "generated_at": datetime.now(timezone.utc).isoformat()
    }
    try:
        await db.report_history.insert_one(record)
    except:
        pass
