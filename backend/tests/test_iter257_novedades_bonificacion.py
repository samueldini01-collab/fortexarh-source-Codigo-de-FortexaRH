"""
Iter 257 — TSS Novedades v5.1 + TSS Bonificación v1.4 native XLSX templates.

Targets:
- GET /api/dgii-reports/native/tss-novedades-v51?year=2026&month=3
- GET /api/dgii-reports/native/tss-bonificacion-v14?year=2026&month=3
- Regression of iter256 endpoints (auto-determinación, IR-4)
"""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get(
    "REACT_APP_BACKEND_URL",
    "https://company-config-debug.preview.emergentagent.com",
).rstrip("/")
ADMIN_EMAIL = "test_refactor@fortexa.com"
ADMIN_PASS = "test123"
YEAR = 2026
MONTH = 3
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(scope="module")
def auth_token():
    r = requests.post(
        f"{BASE_URL}/api/auth/login",
        json={"email": ADMIN_EMAIL, "password": ADMIN_PASS},
        timeout=30,
    )
    if r.status_code != 200:
        pytest.skip(f"Auth failed: {r.status_code} {r.text[:200]}")
    return r.json().get("access_token") or r.json().get("token")


@pytest.fixture(scope="module")
def client(auth_token):
    s = requests.Session()
    s.headers.update({"Authorization": f"Bearer {auth_token}", "Content-Type": "application/json"})
    return s


# ===================== TSS Novedades v5.1 =====================


class TestTssNovedadesV51:
    def test_returns_xlsx_with_correct_mime(self, client):
        r = client.get(
            f"{BASE_URL}/api/dgii-reports/native/tss-novedades-v51",
            params={"year": YEAR, "month": MONTH},
            timeout=60,
        )
        assert r.status_code == 200, r.text[:400]
        assert XLSX_MIME in r.headers.get("Content-Type", ""), r.headers
        assert len(r.content) > 2000, f"file too small: {len(r.content)}"
        assert r.content[:2] == b"PK", "not a real XLSX (no PK ZIP header)"
        cd = r.headers.get("Content-Disposition", "")
        assert "TSS_Novedades_v51" in cd or "novedades" in cd.lower()

    def test_xlsx_structure(self, client):
        r = client.get(
            f"{BASE_URL}/api/dgii-reports/native/tss-novedades-v51",
            params={"year": YEAR, "month": MONTH},
            timeout=60,
        )
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        # Primary sheet
        assert "Plantilla de archivo novedades" in wb.sheetnames, wb.sheetnames
        ws = wb["Plantilla de archivo novedades"]
        # Row 5 title
        assert "Plantilla de Archivo Novedades" in str(ws["B5"].value or "")
        # Row 6 RNC label
        assert "RNC" in str(ws["A6"].value or "")
        # Row 7 Período MMAAAA
        assert "Período" in str(ws["A7"].value or "") or "Periodo" in str(ws["A7"].value or "")
        assert str(ws["B7"].value or "") == f"{MONTH:02d}{YEAR}"
        # Row 9 # Empleados
        assert "Empleados" in str(ws["C9"].value or "")
        # Row 10 group banners
        row10 = "|".join(str(ws.cell(row=10, column=c).value or "") for c in range(1, 30))
        for banner in ("TRABAJADORES", "SDSS", "DGII", "INFOTEP"):
            assert banner in row10, f"missing banner '{banner}' in row 10"
        # Row 11 detail headers
        row11 = " ".join(str(ws.cell(row=11, column=c).value or "") for c in range(2, 25)).lower()
        for keyword in [
            "clave", "tipo", "novedad", "fecha", "inicio", "fin", "doc",
            "documento", "nombres", "apellido", "sexo", "nacimiento",
            "cotizable", "voluntario", "ingreso", "isr", "remuneraciones",
            "agente", "saldo", "regalía", "preaviso", "pensión", "infotep",
        ]:
            assert keyword in row11, f"missing '{keyword}' in row 11 detail headers"

    def test_catalogos_sheet_tnov_codes(self, client):
        r = client.get(
            f"{BASE_URL}/api/dgii-reports/native/tss-novedades-v51",
            params={"year": YEAR, "month": MONTH},
            timeout=60,
        )
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert "Catalogos" in wb.sheetnames
        aux = wb["Catalogos"]
        # Column A holds TNOV codes
        col_a = " ".join(str(aux.cell(row=r, column=1).value or "") for r in range(1, 12))
        for code in ("IN", "SA", "VC", "LV", "LM", "LD", "AD"):
            assert code in col_a, f"missing TNOV code {code} in Catalogos sheet"

    def test_empty_period_still_returns_skeleton(self, client):
        """Novedades endpoint should not 404 — emits empty template skeleton."""
        r = client.get(
            f"{BASE_URL}/api/dgii-reports/native/tss-novedades-v51",
            params={"year": 2099, "month": 1},
            timeout=30,
        )
        # No 404 like the auto/ir4 endpoints; returns the empty skeleton
        assert r.status_code == 200, r.text[:300]
        assert r.content[:2] == b"PK"


# ===================== TSS Bonificación v1.4 =====================


class TestTssBonificacionV14:
    def test_returns_xlsx_with_correct_mime(self, client):
        r = client.get(
            f"{BASE_URL}/api/dgii-reports/native/tss-bonificacion-v14",
            params={"year": YEAR, "month": MONTH},
            timeout=60,
        )
        assert r.status_code == 200, r.text[:400]
        assert XLSX_MIME in r.headers.get("Content-Type", "")
        assert len(r.content) > 1500
        assert r.content[:2] == b"PK"

    def test_xlsx_structure(self, client):
        r = client.get(
            f"{BASE_URL}/api/dgii-reports/native/tss-bonificacion-v14",
            params={"year": YEAR, "month": MONTH},
            timeout=60,
        )
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert "Plantilla de Bonificación" in wb.sheetnames, wb.sheetnames
        ws = wb["Plantilla de Bonificación"]
        # Row 6 title
        assert "Bonificación INFOTEP" in str(ws["A6"].value or "")
        # Row 7 RNC + Ver 1.4
        assert "RNC" in str(ws["A7"].value or "")
        ver_cell = " ".join(str(ws.cell(row=7, column=c).value or "") for c in range(1, 10))
        assert "1.4" in ver_cell, f"Ver. 1.4 not found in row 7: {ver_cell}"
        # Row 8 Período MMAAAA
        assert "Período" in str(ws["A8"].value or "") or "Periodo" in str(ws["A8"].value or "")
        assert str(ws["B8"].value or "") == f"{MONTH:02d}{YEAR}"
        # Row 10 # Empleados
        assert "Empleados" in str(ws["C10"].value or "")
        # Row 11 INFOTEP group banner
        row11 = " ".join(str(ws.cell(row=11, column=c).value or "") for c in range(1, 12))
        assert "INFOTEP" in row11
        # Row 12 headers
        row12 = " ".join(str(ws.cell(row=12, column=c).value or "") for c in range(2, 10)).lower()
        for keyword in [
            "tipo doc", "número doc", "nombres", "1er. apellido", "2do. apellido",
            "sexo", "fecha nacimiento", "monto bonificación",
        ]:
            assert keyword in row12, f"missing '{keyword}' in row 12 headers"

    def test_empty_bonuses_still_returns_skeleton(self, client):
        """Bonificación should return template skeleton (HTTP 200) even with no bonuses."""
        r = client.get(
            f"{BASE_URL}/api/dgii-reports/native/tss-bonificacion-v14",
            params={"year": 2099, "month": 1},
            timeout=30,
        )
        assert r.status_code == 200, r.text[:300]
        assert XLSX_MIME in r.headers.get("Content-Type", "")
        assert r.content[:2] == b"PK"


# ===================== Regression: iter256 endpoints =====================


class TestIter256Regression:
    def test_tss_auto_v53_still_works(self, client):
        r = client.get(
            f"{BASE_URL}/api/dgii-reports/native/tss-autodeterminacion-v53",
            params={"year": YEAR, "month": MONTH},
            timeout=60,
        )
        assert r.status_code == 200
        assert XLSX_MIME in r.headers.get("Content-Type", "")
        assert r.content[:2] == b"PK"

    def test_ir4_official_still_works(self, client):
        r = client.get(
            f"{BASE_URL}/api/dgii-reports/native/ir4-official",
            params={"year": YEAR, "month": MONTH},
            timeout=60,
        )
        assert r.status_code == 200
        assert XLSX_MIME in r.headers.get("Content-Type", "")
        assert r.content[:2] == b"PK"
