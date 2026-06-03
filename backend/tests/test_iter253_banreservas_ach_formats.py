"""
Iter 253 — Banreservas ACH dual format (TXT + XLSX) backend tests.

Covers:
  - GET /api/bank-files/banks returns Banreservas with formats=['txt','xlsx']
  - GET /api/bank-files/generate/{period}/banreservas (default = txt)
  - GET /api/bank-files/generate/{period}/banreservas?format=xlsx returns Excel
  - Sanity: txt vs xlsx record_count match for same period
  - Other banks (popular, bhd) ignore format=xlsx -> return TXT
"""
import io
import os
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/") or "https://company-config-debug.preview.emergentagent.com"
ADMIN_EMAIL = "test_refactor@fortexa.com"
ADMIN_PASS = "test123"
PERIOD_ID = "period_42a51c051530"


@pytest.fixture(scope="module")
def auth_headers():
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login",
               json={"email": ADMIN_EMAIL, "password": ADMIN_PASS},
               timeout=30)
    if r.status_code != 200:
        pytest.skip(f"Login failed: {r.status_code} {r.text[:120]}")
    token = r.json().get("access_token") or r.json().get("token")
    if not token:
        pytest.skip(f"No token in login response: {r.json()}")
    return {"Authorization": f"Bearer {token}"}


# --- /banks listing ---
def test_banks_listing_includes_banreservas_dual_format(auth_headers):
    r = requests.get(f"{BASE_URL}/api/bank-files/banks", headers=auth_headers, timeout=15)
    assert r.status_code == 200, r.text
    banks = r.json()
    assert isinstance(banks, list) and len(banks) >= 3
    br = next((b for b in banks if b["id"] == "banreservas"), None)
    assert br is not None, "Banreservas missing from /banks"
    assert "formats" in br, "formats key missing"
    assert set(br["formats"]) == {"txt", "xlsx"}, br["formats"]
    # description should mention both formats
    desc = (br.get("description") or "").lower()
    assert "txt" in desc and ("xlsx" in desc or "excel" in desc), br.get("description")

    # popular / bhd should be txt-only
    for bid in ("popular", "bhd"):
        b = next((x for x in banks if x["id"] == bid), None)
        assert b is not None
        assert b.get("formats") == ["txt"], f"{bid} formats={b.get('formats')}"


# --- TXT (default) generation ---
def test_banreservas_txt_default(auth_headers):
    r = requests.get(
        f"{BASE_URL}/api/bank-files/generate/{PERIOD_ID}/banreservas",
        headers=auth_headers, timeout=30,
    )
    assert r.status_code == 200, r.text
    ctype = r.headers.get("Content-Type", "")
    assert "text/plain" in ctype, ctype
    cd = r.headers.get("Content-Disposition", "")
    assert ".txt" in cd, cd
    body = r.text.strip().splitlines()
    assert len(body) >= 1, "Expected at least 1 ACH line"
    # Each line must have exactly 8 comma-separated fields
    for line in body:
        parts = line.split(",")
        assert len(parts) == 8, f"line has {len(parts)} parts: {line}"
        assert parts[0] in ("CC", "CA"), parts[0]
        assert parts[1] in ("DOP", "USD"), parts[1]
        assert parts[3] in ("CC", "CA"), parts[3]
        # amount must be float-castable with 2 decimals
        float(parts[6])
    rec_count = int(r.headers.get("X-Record-Count", "0"))
    assert rec_count == len(body)
    return rec_count, body


# --- XLSX generation ---
def test_banreservas_xlsx_format(auth_headers):
    r = requests.get(
        f"{BASE_URL}/api/bank-files/generate/{PERIOD_ID}/banreservas",
        params={"format": "xlsx"},
        headers=auth_headers, timeout=30,
    )
    assert r.status_code == 200, r.text[:300]
    ctype = r.headers.get("Content-Type", "")
    assert "spreadsheetml.sheet" in ctype, ctype
    cd = r.headers.get("Content-Disposition", "")
    assert ".xlsx" in cd, cd
    # Try to open with openpyxl
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert "CONVERTIR NOMINA" in wb.sheetnames, wb.sheetnames
    ws = wb["CONVERTIR NOMINA"]
    # Row 1 col A label / col E value
    assert ws.cell(row=1, column=1).value == "TIPO CUENTA EMPRESA"
    assert ws.cell(row=2, column=1).value == "NUMERO DE CUENTA EMPRESA"
    assert ws.cell(row=3, column=1).value == "MONEDA A PAGAR"
    assert ws.cell(row=4, column=1).value == "TOTAL DE NÓMINA"
    assert ws.cell(row=5, column=1).value == "CANTIDAD DE EMPLEADOS"
    # meta values should be present (non-empty) in col E
    for row in (1, 2, 3, 4, 5):
        assert ws.cell(row=row, column=5).value not in (None, ""), f"row {row} col E empty"
    # Row 8 has 9 headers
    header_row = [ws.cell(row=8, column=c).value for c in range(1, 10)]
    assert all(h for h in header_row), f"Row 8 headers incomplete: {header_row}"
    assert header_row[0].lower().startswith("nombre"), header_row[0]
    # Data starts at row 9
    employees_in_meta = int(ws.cell(row=5, column=5).value)
    data_rows = 0
    for row in ws.iter_rows(min_row=9, values_only=True):
        if any(c is not None and str(c).strip() != "" for c in row):
            data_rows += 1
    assert data_rows == employees_in_meta, f"data rows={data_rows} meta count={employees_in_meta}"
    assert int(r.headers.get("X-Record-Count", "0")) == data_rows


# --- sanity: same record count for txt vs xlsx ---
def test_txt_and_xlsx_same_record_count(auth_headers):
    r_txt = requests.get(
        f"{BASE_URL}/api/bank-files/generate/{PERIOD_ID}/banreservas",
        headers=auth_headers, timeout=30,
    )
    r_xlsx = requests.get(
        f"{BASE_URL}/api/bank-files/generate/{PERIOD_ID}/banreservas",
        params={"format": "xlsx"}, headers=auth_headers, timeout=30,
    )
    assert r_txt.status_code == 200 and r_xlsx.status_code == 200
    assert r_txt.headers.get("X-Record-Count") == r_xlsx.headers.get("X-Record-Count")


# --- format=xlsx ignored for other banks ---
@pytest.mark.parametrize("bank_id", ["popular", "bhd"])
def test_other_banks_ignore_xlsx_param(auth_headers, bank_id):
    r = requests.get(
        f"{BASE_URL}/api/bank-files/generate/{PERIOD_ID}/{bank_id}",
        params={"format": "xlsx"},
        headers=auth_headers, timeout=30,
    )
    # Should still succeed; should return TXT (or 400 if no bank data — we accept either non-xlsx outcome)
    if r.status_code != 200:
        # Acceptable when employees lack accounts for this bank fmt; just ensure it's not xlsx
        assert "spreadsheetml" not in r.headers.get("Content-Type", "")
        return
    ctype = r.headers.get("Content-Type", "")
    assert "spreadsheetml" not in ctype, f"{bank_id} returned xlsx unexpectedly"
    cd = r.headers.get("Content-Disposition", "")
    assert ".txt" in cd, cd
