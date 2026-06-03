"""
Iter 256 — Native DGII/TSS XLSX template generators + bank-info regression.

Targets:
- /api/dgii-reports/native/tss-autodeterminacion-v53
- /api/dgii-reports/native/ir4-official
- Legacy /api/dgii-reports/monthly/* (backward-compat)
- /api/employees/{id}/bank-info PATCH (regression)
"""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://company-config-debug.preview.emergentagent.com").rstrip("/")
ADMIN_EMAIL = "test_refactor@fortexa.com"
ADMIN_PASS = "test123"
YEAR = 2026
MONTH = 3

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(scope="module")
def auth_token():
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASS}, timeout=30)
    if r.status_code != 200:
        pytest.skip(f"Auth failed: {r.status_code} {r.text[:200]}")
    return r.json().get("access_token") or r.json().get("token")


@pytest.fixture(scope="module")
def client(auth_token):
    s = requests.Session()
    s.headers.update({"Authorization": f"Bearer {auth_token}", "Content-Type": "application/json"})
    return s


# ===================== TSS Autodeterminación v5.3 =====================

class TestTssAutoNative:
    def test_returns_xlsx_with_correct_mime(self, client):
        r = client.get(f"{BASE_URL}/api/dgii-reports/native/tss-autodeterminacion-v53",
                       params={"year": YEAR, "month": MONTH}, timeout=60)
        assert r.status_code == 200, r.text[:400]
        assert XLSX_MIME in r.headers.get("Content-Type", ""), r.headers
        assert len(r.content) > 2000, f"file too small: {len(r.content)}"
        assert r.content[:2] == b"PK", "not a real XLSX (no PK ZIP header)"

    def test_xlsx_structure(self, client):
        r = client.get(f"{BASE_URL}/api/dgii-reports/native/tss-autodeterminacion-v53",
                       params={"year": YEAR, "month": MONTH}, timeout=60)
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert "Plantilla de Autodeterminación" in wb.sheetnames
        ws = wb["Plantilla de Autodeterminación"]
        # Header expectations per request
        assert "AutoDeterminación" in str(ws["A5"].value or "")
        assert "Tipo" in str(ws["A6"].value or "")
        assert str(ws["B6"].value or "").upper() == "AM"
        assert "RNC" in str(ws["A7"].value or "")
        assert "Período" in str(ws["A8"].value or "") or "Periodo" in str(ws["A8"].value or "")
        # Period MMAAAA
        assert str(ws["B8"].value or "") == f"{MONTH:02d}{YEAR}"
        # # Empleados in row 10
        assert "Empleados" in str(ws["C10"].value or "")
        assert isinstance(ws["E10"].value, int) and ws["E10"].value >= 1
        # Group headers row 11
        row11 = [str(ws.cell(row=11, column=c).value or "") for c in range(1, 25)]
        joined = "|".join(row11)
        assert "TRABAJADORES" in joined
        assert "SDSS" in joined
        assert "DGII" in joined
        assert "INFOTEP" in joined
        # Detail headers row 12 from col B..U (2..21)
        row12 = [str(ws.cell(row=12, column=c).value or "") for c in range(2, 22)]
        joined12 = " ".join(row12).lower()
        for expected in ["clave", "tipo", "documento", "nombres", "apellido", "sexo",
                         "nacimiento", "cotizable", "voluntario", "isr", "ingreso",
                         "remuneraciones", "agente", "saldo", "regalía", "preaviso",
                         "pensión", "infotep"]:
            assert expected in joined12, f"missing '{expected}' in row 12 headers"
        # Data starts row 13
        assert ws.cell(row=13, column=2).value is not None, "no data on row 13"

    def test_404_when_no_data(self, client):
        r = client.get(f"{BASE_URL}/api/dgii-reports/native/tss-autodeterminacion-v53",
                       params={"year": 2099, "month": 1}, timeout=30)
        assert r.status_code == 404, r.text[:300]


# ===================== IR-4 Oficial =====================

class TestIr4Native:
    def test_returns_xlsx_with_correct_mime(self, client):
        r = client.get(f"{BASE_URL}/api/dgii-reports/native/ir4-official",
                       params={"year": YEAR, "month": MONTH}, timeout=60)
        assert r.status_code == 200, r.text[:400]
        assert XLSX_MIME in r.headers.get("Content-Type", "")
        assert len(r.content) > 2000
        assert r.content[:2] == b"PK"

    def test_xlsx_structure(self, client):
        r = client.get(f"{BASE_URL}/api/dgii-reports/native/ir4-official",
                       params={"year": YEAR, "month": MONTH}, timeout=60)
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        assert "DIRECCIÓN GENERAL DE IMPUESTOS INTERNOS" in str(ws["D2"].value or "")
        assert "IR-4" in str(ws["D3"].value or "")
        assert "AGENTE DE RETENCIÓN" in str(ws["A5"].value or "")
        assert "RNC" in str(ws["G5"].value or "")
        assert "DESDE" in str(ws["A6"].value or "")
        assert "HASTA" in str(ws["D6"].value or "")
        # Section banners
        assert "ASALARIADO" in str(ws["A8"].value or "")
        assert "REMUNERACIONES" in str(ws["D8"].value or "")
        # Column headers row 10 (A..M = 1..13)
        headers = [str(ws.cell(row=10, column=c).value or "") for c in range(1, 14)]
        joined = " ".join(headers).lower()
        assert "no." in joined
        assert "apellidos" in joined and "nombres" in joined
        assert "cédula" in joined or "cedula" in joined or "rnc" in joined
        assert "sueldos" in joined
        # Totals row at the bottom
        n_data = 0
        for row in range(11, ws.max_row + 1):
            c2 = ws.cell(row=row, column=2).value
            if c2 == "TOTALES":
                # totals row exists
                assert isinstance(ws.cell(row=row, column=4).value, (int, float))
                break
            elif c2:
                n_data += 1
        else:
            pytest.fail("TOTALES row not found")
        assert n_data >= 1

    def test_404_when_no_data(self, client):
        r = client.get(f"{BASE_URL}/api/dgii-reports/native/ir4-official",
                       params={"year": 2099, "month": 1}, timeout=30)
        assert r.status_code == 404


# ===================== Legacy regression =====================

class TestLegacyDgiiMonthly:
    def test_legacy_ir4(self, client):
        r = client.get(f"{BASE_URL}/api/dgii-reports/monthly/ir4",
                       params={"year": YEAR, "month": MONTH}, timeout=60)
        assert r.status_code == 200, r.text[:300]
        # Legacy is tab-delimited
        ct = r.headers.get("Content-Type", "")
        assert XLSX_MIME not in ct, f"expected tab-delimited, got {ct}"

    def test_legacy_tss_auto(self, client):
        r = client.get(f"{BASE_URL}/api/dgii-reports/monthly/tss-autodeterminacion",
                       params={"year": YEAR, "month": MONTH}, timeout=60)
        assert r.status_code == 200, r.text[:300]
        ct = r.headers.get("Content-Type", "")
        assert XLSX_MIME not in ct


# ===================== Bank-info PATCH regression =====================

class TestBankInfoPatch:
    def test_patch_bank_info(self, client):
        # Pick first employee
        r = client.get(f"{BASE_URL}/api/employees", timeout=30)
        assert r.status_code == 200
        emps = r.json() if isinstance(r.json(), list) else r.json().get("employees", [])
        assert len(emps) > 0
        emp_id = emps[0].get("employee_id") or emps[0].get("id")
        assert emp_id

        payload = {
            "bank_name": "Banco de Reservas (Banreservas)",
            "bank_code": "002",
            "account_number": "9999999999",
            "account_type": "checking",
        }
        r = client.patch(f"{BASE_URL}/api/employees/{emp_id}/bank-info", json=payload, timeout=30)
        assert r.status_code in (200, 204), r.text[:300]
        # Verify
        rv = client.get(f"{BASE_URL}/api/employees/{emp_id}", timeout=30)
        assert rv.status_code == 200
        data = rv.json()
        assert data.get("bank_name") == "Banco de Reservas (Banreservas)"
